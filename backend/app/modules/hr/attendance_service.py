import logging
from datetime import date, timedelta
from typing import Optional

from sqlalchemy import func, asc, desc
from sqlalchemy.orm import Session, selectinload

from app.core.pagination import legacy_list, page_of, wants_page

logger = logging.getLogger("zoiko")

from app.modules.hr.models import (
    AttendanceRecord, Shift, ShiftRoster, Holiday, Employee, Department,
    AttendanceStatus, ShiftType, LeaveRequest, LeaveBalance, RequestStatus,
    LeaveType, UserRole, Organization,
)
from app.modules.hr.schemas import (
    AttendanceCreate, AttendanceUpdate,
    ShiftCreate, ShiftUpdate,
    ShiftRosterCreate,
    HolidayCreate, HolidayUpdate,
    SuccessResponse,
)
from app.core.exceptions import NotFoundException, BadRequestException
from app.core.sanitize import sanitize_dict


# ── DASHBOARD ─────────────────────────────────────────────────────────────────

STANDARD_WORKDAY_HOURS = 8.0
ATTENDED_STATUSES = (AttendanceStatus.PRESENT, AttendanceStatus.LATE, AttendanceStatus.REMOTE, AttendanceStatus.HALF_DAY)


def _record_hours(rec) -> Optional[float]:
    if rec.total_hours is not None:
        return float(rec.total_hours)
    if rec.check_in and rec.check_out and rec.check_out > rec.check_in:
        return (rec.check_out - rec.check_in).total_seconds() / 3600
    return None


def _day_snapshot(db: Session, day: date, organization_id: Optional[int], department: Optional[str], workforce_ids: set) -> dict:
    """Head-counts for one day. Every figure counts distinct people from the current workforce, so
    "present" can never exceed the workforce (the page used to show 2/1 and 2/0)."""
    query = db.query(AttendanceRecord).filter(AttendanceRecord.date == day, AttendanceRecord.is_deleted.isnot(True))
    if organization_id:
        query = query.filter(AttendanceRecord.organization_id == organization_id)
    records = [r for r in query.all() if r.employee_id in workforce_ids]

    by_status: dict = {}
    for r in records:
        by_status.setdefault(r.status, set()).add(r.employee_id)
    attended = set().union(*(by_status.get(s, set()) for s in ATTENDED_STATUSES))
    return {"records": records, "by_status": by_status, "attended": attended}


def _change(today_value, yesterday_value) -> Optional[float]:
    """Percentage change against yesterday; None when yesterday has nothing to compare with."""
    if not yesterday_value:
        return None
    return round((today_value - yesterday_value) / yesterday_value * 100, 1)


def get_attendance_dashboard(db: Session, organization_id: Optional[int] = None, department: Optional[str] = None) -> dict:
    """Live attendance figures for today, all computed from real records for this organization."""
    from app.modules.employee.models import EmployeeStatus

    today = date.today()
    yesterday = today - timedelta(days=1)

    emp_query = db.query(Employee).filter(Employee.is_active.is_(True), Employee.status.in_([EmployeeStatus.ACTIVE, EmployeeStatus.ON_LEAVE]))
    if organization_id:
        emp_query = emp_query.filter(Employee.organization_id == organization_id)
    everyone = emp_query.all()  # admins and managers check in too, so they belong in the headcount
    dept_names = {d.id: d.name for d in db.query(Department).filter(Department.organization_id == organization_id).all()} if organization_id else {d.id: d.name for d in db.query(Department).all()}
    departments_available = sorted({dept_names[e.department_id] for e in everyone if e.department_id in dept_names})
    workforce = [e for e in everyone if not department or dept_names.get(e.department_id) == department]
    workforce_ids = {e.id for e in workforce}
    total_employees = len(workforce)

    snap = _day_snapshot(db, today, organization_id, department, workforce_ids)
    prev = _day_snapshot(db, yesterday, organization_id, department, workforce_ids)
    present = len(snap["attended"])
    absent = len(snap["by_status"].get(AttendanceStatus.ABSENT, set()))
    on_leave = len(snap["by_status"].get(AttendanceStatus.ON_LEAVE, set()))
    remote = len(snap["by_status"].get(AttendanceStatus.REMOTE, set()))
    late = len(snap["by_status"].get(AttendanceStatus.LATE, set()))

    shift_by_employee = {}
    shift_rows = (
        db.query(ShiftRoster.employee_id, Shift)
        .join(Shift, Shift.id == ShiftRoster.shift_id)
        .filter(ShiftRoster.date == today, ShiftRoster.is_active.is_(True))
    )
    if organization_id:
        shift_rows = shift_rows.filter(Shift.organization_id == organization_id)
    for emp_id, shift in shift_rows.all():
        if emp_id in workforce_ids:
            shift_by_employee[emp_id] = shift

    def _early(rec) -> bool:
        shift = shift_by_employee.get(rec.employee_id)
        if not (shift and rec.check_out and shift.end_time):
            return False
        try:
            hh, mm = (int(x) for x in shift.end_time.split(":")[:2])
        except ValueError:
            return False
        return (rec.check_out.hour, rec.check_out.minute) < (hh, mm)

    early_departures = len({r.employee_id for r in snap["records"] if r.status in ATTENDED_STATUSES and _early(r)})

    per_person: dict = {}  # one figure per person, so a duplicated record cannot double their hours
    for r in snap["records"]:
        h = _record_hours(r) if r.check_out else None
        if h is not None:
            per_person[r.employee_id] = max(h, per_person.get(r.employee_id, 0.0))
    worked = list(per_person.values())
    avg_hours = round(sum(worked) / len(worked), 2) if worked else 0.0
    overtime = round(sum(max(h - STANDARD_WORKDAY_HOURS, 0) for h in worked), 2)

    # per department: how many of its people are in today
    dept_rows = []
    for name in sorted({dept_names[e.department_id] for e in workforce if e.department_id in dept_names}):
        members = {e.id for e in workforce if dept_names.get(e.department_id) == name}
        dept_rows.append({"department": name, "present": len(members & snap["attended"]), "total": len(members), "count": len(members & snap["attended"])})

    shift_counts: dict = {}
    for shift in shift_by_employee.values():
        shift_counts[shift.name] = shift_counts.get(shift.name, 0) + 1
    shift_dist = [{"shift": n, "count": c} for n, c in sorted(shift_counts.items())]

    trend = []
    for offset in range(6, -1, -1):
        day = today - timedelta(days=offset)
        s = snap if offset == 0 else (prev if offset == 1 else _day_snapshot(db, day, organization_id, department, workforce_ids))
        trend.append({
            "label": day.strftime("%a"), "date": day.isoformat(),
            "present": len(s["attended"]), "absent": len(s["by_status"].get(AttendanceStatus.ABSENT, set())),
        })

    prev_status = prev["by_status"]
    percentage = round(present / total_employees * 100, 2) if total_employees else 0.0
    changes = {
        "present_today": _change(present, len(prev["attended"])),
        "absent_today": _change(absent, len(prev_status.get(AttendanceStatus.ABSENT, set()))),
        "late_arrivals": _change(late, len(prev_status.get(AttendanceStatus.LATE, set()))),
        "on_leave": _change(on_leave, len(prev_status.get(AttendanceStatus.ON_LEAVE, set()))),
        "remote": _change(remote, len(prev_status.get(AttendanceStatus.REMOTE, set()))),
    }
    return {
        "present_today": present,
        "absent_today": absent,
        "late_arrivals": late,
        "early_departures": early_departures,
        "on_leave": on_leave,
        "on_leave_count": on_leave,
        "remote": remote,
        "remote_count": remote,
        "overtime": overtime,
        "overtime_count": overtime,
        "attendance_percentage": percentage,
        "attendance_rate": percentage,
        "avg_working_hours": avg_hours,
        "total_employees": total_employees,
        "department_attendance": dept_rows,
        "department_breakdown": dept_rows,
        "shift_distribution": shift_dist,
        "shift_utilization": shift_dist,
        "attendance_trend": trend,
        "changes": changes,
        "departments": departments_available,
        "as_of": today.isoformat(),
    }


# ── ATTENDANCE RECORDS CRUD ──────────────────────────────────────────────────

SORTABLE_FIELDS_RECORDS = {
    "id": AttendanceRecord.id,
    "employee_id": AttendanceRecord.employee_id,
    "date": AttendanceRecord.date,
    "status": AttendanceRecord.status,
    "check_in": AttendanceRecord.check_in,
    "check_out": AttendanceRecord.check_out,
    "created_at": AttendanceRecord.created_at,
}


def _get_records_query(
    db: Session, search=None, status=None, department=None,
    date_from=None, date_to=None, employee_id=None, organization_id=None,
):
    query = db.query(AttendanceRecord)

    if organization_id:
        query = query.filter(AttendanceRecord.organization_id == organization_id)

    if search:
        stmt = f"%{search}%"
        query = query.join(Employee, AttendanceRecord.employee_id == Employee.id).filter(
            (Employee.first_name.ilike(stmt)) |
            (Employee.last_name.ilike(stmt)) |
            (Employee.employee_code.ilike(stmt))
        )

    if status:
        query = query.filter(AttendanceRecord.status == status)

    if department:
        query = query.join(Employee, AttendanceRecord.employee_id == Employee.id).filter(
            Employee.department_id == Department.id,
            Department.name == department,
        )

    if date_from:
        query = query.filter(AttendanceRecord.date >= date_from)

    if date_to:
        query = query.filter(AttendanceRecord.date <= date_to)

    if employee_id:
        query = query.filter(AttendanceRecord.employee_id == employee_id)

    return query


def get_all_attendance_records(db: Session, organization_id: Optional[int] = None, employee_id: Optional[int] = None,
                               page: Optional[int] = None, per_page: Optional[int] = None):
    """Every record, newest first: a plain list (capped) without page/per_page, else a page dict."""
    query = db.query(AttendanceRecord)
    if organization_id:
        query = query.filter(AttendanceRecord.organization_id == organization_id)
    if employee_id:
        query = query.filter(AttendanceRecord.employee_id == employee_id)
    # employee (and its department) were lazy-loaded per record: one query per distinct employee. Batch them.
    loaded = query.options(selectinload(AttendanceRecord.employee).joinedload(Employee.department))
    if wants_page(page, per_page):
        # pages need a total order: id breaks ties between records of the same date
        paged = page_of(loaded.order_by(AttendanceRecord.date.desc(), AttendanceRecord.id.desc()), page, per_page)
        records = paged["items"]
    else:
        paged = None
        records = legacy_list(loaded.order_by(AttendanceRecord.date.desc()), "/hr/attendance")
    items = []
    for r in records:
        items.append({
            "id": r.id,
            "employee_id": r.employee_id,
            "date": r.date,
            "status": r.status.value if hasattr(r.status, 'value') else r.status,
            "check_in": r.check_in,
            "check_out": r.check_out,
            "notes": r.notes,
            "created_at": r.created_at,
            "employee_name": r.employee.full_name if r.employee else None,
            "department": r.employee.department.name if r.employee and r.employee.department else None,
        })
    if paged:
        return {**paged, "items": items}
    return items


def get_attendance_records(
    db: Session,
    page: int = 1,
    per_page: int = 20,
    search: Optional[str] = None,
    status: Optional[AttendanceStatus] = None,
    department: Optional[str] = None,
    date_from: Optional[date] = None,
    date_to: Optional[date] = None,
    employee_id: Optional[int] = None,
    sort_by: Optional[str] = "date",
    sort_order: Optional[str] = "desc",
    organization_id: Optional[int] = None,
) -> dict:
    per_page = min(per_page, 100)
    query = _get_records_query(db, search, status, department, date_from, date_to, employee_id, organization_id)
    total = query.count()

    sort_col = SORTABLE_FIELDS_RECORDS.get(sort_by, AttendanceRecord.date)
    sort_fn = desc if sort_order == "desc" else asc
    records = query.options(selectinload(AttendanceRecord.employee)).order_by(sort_fn(sort_col)).offset((page - 1) * per_page).limit(per_page).all()

    items = []
    for r in records:
        items.append({
            "id": r.id,
            "employee_id": r.employee_id,
            "date": r.date,
            "status": r.status,
            "check_in": r.check_in,
            "check_out": r.check_out,
            "notes": r.notes,
            "created_at": r.created_at,
            "employee_name": r.employee.full_name if r.employee else None,
        })

    return {"total": total, "page": page, "per_page": per_page, "items": items}


def get_attendance_record_by_id(db: Session, record_id: int, organization_id: Optional[int] = None) -> AttendanceRecord:
    query = db.query(AttendanceRecord).filter(AttendanceRecord.id == record_id)
    if organization_id:
        query = query.filter(AttendanceRecord.organization_id == organization_id)
    record = query.first()
    if not record:
        raise NotFoundException("AttendanceRecord", record_id)
    return record


def create_attendance_record(db: Session, data: AttendanceCreate, created_by: int = None, organization_id: int = None) -> AttendanceRecord:
    raw = data.model_dump()
    safe = sanitize_dict(raw)
    record = AttendanceRecord(**safe)
    if organization_id:
        record.organization_id = organization_id
    db.add(record)
    db.commit()
    db.refresh(record)
    return record


def update_attendance_record(db: Session, record_id: int, data: AttendanceUpdate, organization_id: Optional[int] = None) -> AttendanceRecord:
    record = get_attendance_record_by_id(db, record_id, organization_id)
    update_data = sanitize_dict(data.model_dump(exclude_unset=True))
    for field, value in update_data.items():
        setattr(record, field, value)
    db.commit()
    db.refresh(record)
    return record


def delete_attendance_record(db: Session, record_id: int, organization_id: Optional[int] = None) -> None:
    record = get_attendance_record_by_id(db, record_id, organization_id)
    db.delete(record)
    db.commit()


# ── SHIFTS ───────────────────────────────────────────────────────────────────

def get_shifts(db: Session, organization_id: Optional[int] = None) -> list[Shift]:
    query = db.query(Shift)
    if organization_id:
        query = query.filter(Shift.organization_id == organization_id)
    return query.order_by(Shift.name).all()


def create_shift(db: Session, data: ShiftCreate, created_by: int = None, organization_id: int = None) -> Shift:
    shift = Shift(**data.model_dump(), created_by=created_by)
    if organization_id:
        shift.organization_id = organization_id
    db.add(shift)
    db.commit()
    db.refresh(shift)
    return shift


def get_shift_by_id(db: Session, shift_id: int, organization_id: Optional[int] = None) -> Shift:
    query = db.query(Shift).filter(Shift.id == shift_id)
    if organization_id:
        query = query.filter(Shift.organization_id == organization_id)
    shift = query.first()
    if not shift:
        raise NotFoundException("Shift", shift_id)
    return shift


def update_shift(db: Session, shift_id: int, data: ShiftUpdate, organization_id: Optional[int] = None) -> Shift:
    shift = get_shift_by_id(db, shift_id, organization_id)
    update_data = sanitize_dict(data.model_dump(exclude_unset=True))
    for field, value in update_data.items():
        setattr(shift, field, value)
    db.commit()
    db.refresh(shift)
    return shift


def delete_shift(db: Session, shift_id: int, organization_id: Optional[int] = None) -> None:
    shift = get_shift_by_id(db, shift_id, organization_id)
    db.delete(shift)
    db.commit()


# ── SHIFT ROSTERS ────────────────────────────────────────────────────────────

def get_shift_rosters(
    db: Session, page: int = 1, per_page: int = 20,
    date_filter: Optional[date] = None,
    employee_id: Optional[int] = None,
    shift_id: Optional[int] = None,
    organization_id: Optional[int] = None,
) -> dict:
    per_page = min(per_page, 100)
    query = db.query(ShiftRoster, Shift, Employee).outerjoin(Shift, ShiftRoster.shift_id == Shift.id).outerjoin(Employee, ShiftRoster.employee_id == Employee.id)
    if organization_id:
        query = query.filter(Employee.organization_id == organization_id)
    if date_filter:
        query = query.filter(ShiftRoster.date == date_filter)
    if employee_id:
        query = query.filter(ShiftRoster.employee_id == employee_id)
    if shift_id:
        query = query.filter(ShiftRoster.shift_id == shift_id)
    total = query.count()
    rows = query.order_by(ShiftRoster.date.desc()).offset((page - 1) * per_page).limit(per_page).all()
    result = []
    for roster, shift, employee in rows:
        result.append({
            "id": roster.id,
            "employee_id": roster.employee_id,
            "shift_id": roster.shift_id,
            "date": roster.date,
            "is_active": roster.is_active,
            "assigned_by": roster.assigned_by,
            "created_at": roster.created_at,
            "updated_at": roster.updated_at,
            "employee_name": employee.full_name if employee else None,
            "shift_name": shift.name if shift else None,
        })
    return {"total": total, "page": page, "per_page": per_page, "items": result}


def create_shift_roster(db: Session, data: ShiftRosterCreate, assigned_by: int = None) -> ShiftRoster:
    roster = ShiftRoster(**data.model_dump(), assigned_by=assigned_by)
    db.add(roster)
    db.commit()
    db.refresh(roster)
    return roster


def delete_shift_roster(db: Session, roster_id: int, organization_id: Optional[int] = None) -> None:
    query = db.query(ShiftRoster).filter(ShiftRoster.id == roster_id)
    if organization_id:
        query = query.join(Employee, ShiftRoster.employee_id == Employee.id).filter(Employee.organization_id == organization_id)
    roster = query.first()
    if not roster:
        raise NotFoundException("ShiftRoster", roster_id)
    db.delete(roster)
    db.commit()


# ── HOLIDAYS ─────────────────────────────────────────────────────────────────

def get_holidays(db: Session, organization_id: Optional[int] = None) -> list[Holiday]:
    query = db.query(Holiday)
    if organization_id:
        query = query.filter(Holiday.organization_id == organization_id)
    return query.order_by(Holiday.date).all()


def create_holiday(db: Session, data: HolidayCreate, created_by: int = None, organization_id: int = None) -> Holiday:
    dup = db.query(Holiday).filter(Holiday.date == data.date, Holiday.name.ilike(data.name))
    if organization_id:
        dup = dup.filter(Holiday.organization_id == organization_id)
    if dup.first():
        raise BadRequestException(f"Holiday '{data.name}' on {data.date} already exists")
    holiday = Holiday(**data.model_dump(), created_by=created_by)
    if organization_id:
        holiday.organization_id = organization_id
    db.add(holiday)
    db.commit()
    db.refresh(holiday)
    return holiday


def get_holiday_by_id(db: Session, holiday_id: int, organization_id: Optional[int] = None) -> Holiday:
    query = db.query(Holiday).filter(Holiday.id == holiday_id)
    if organization_id:
        query = query.filter(Holiday.organization_id == organization_id)
    holiday = query.first()
    if not holiday:
        raise NotFoundException("Holiday", holiday_id)
    return holiday


def update_holiday(db: Session, holiday_id: int, data: HolidayUpdate, organization_id: Optional[int] = None) -> Holiday:
    holiday = get_holiday_by_id(db, holiday_id, organization_id)
    update_data = sanitize_dict(data.model_dump(exclude_unset=True))
    for field, value in update_data.items():
        setattr(holiday, field, value)
    db.commit()
    db.refresh(holiday)
    return holiday


def delete_holiday(db: Session, holiday_id: int, organization_id: Optional[int] = None) -> None:
    holiday = get_holiday_by_id(db, holiday_id, organization_id)
    db.delete(holiday)
    db.commit()


HOLIDAY_TYPES = {"public": "Public", "company": "Company", "optional": "Optional"}
MAX_HOLIDAY_IMPORT_ROWS = 1000


def _parse_holiday_date(value):
    """ISO first, then the day-first formats people type into spreadsheets."""
    from datetime import datetime as _dt

    if isinstance(value, _dt):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value or "").strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y", "%Y/%m/%d"):
        try:
            return _dt.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def _truthy(value) -> bool:
    return str(value).strip().lower() in ("1", "true", "yes", "y", "recurring")


def import_holidays(db: Session, holidays: list[dict], created_by: int = None, organization_id: int = None) -> dict:
    """Import holidays one row at a time and say what happened to each row.

    Returns {"total", "imported", "duplicates", "errors": [{"row", "name", "error"}]}. Bad rows are
    reported, never silently dropped, and never stop the good rows from being saved. Duplicates are
    judged within this organization only (same date and name, ignoring case)."""
    if len(holidays) > MAX_HOLIDAY_IMPORT_ROWS:
        raise BadRequestException(f"Import at most {MAX_HOLIDAY_IMPORT_ROWS} holidays at a time.")
    if not holidays:
        raise BadRequestException("There are no holidays to import.")

    imported, duplicates, errors = 0, 0, []
    seen = set()
    for index, raw in enumerate(holidays, start=1):
        if not isinstance(raw, dict):
            errors.append({"row": index, "name": "", "error": "Each holiday must be an object with a name and a date."})
            continue
        safe = sanitize_dict(raw)
        name = str(safe.get("name") or "").strip()
        if not name:
            errors.append({"row": index, "name": "", "error": "Name is required."})
            continue
        if len(name) > 150:
            errors.append({"row": index, "name": name[:40], "error": "Name is longer than 150 characters."})
            continue
        h_date = _parse_holiday_date(safe.get("date"))
        if h_date is None:
            errors.append({"row": index, "name": name, "error": f"Date '{safe.get('date') or ''}' is missing or not a valid date (use YYYY-MM-DD)."})
            continue
        type_key = str(safe.get("type") or "public").strip().lower()
        if type_key not in HOLIDAY_TYPES:
            errors.append({"row": index, "name": name, "error": f"Type '{safe.get('type')}' is not one of Public, Company or Optional."})
            continue

        key = (h_date, name.lower())
        existing = db.query(Holiday).filter(Holiday.date == h_date, Holiday.name.ilike(name))
        if organization_id is not None:
            existing = existing.filter(Holiday.organization_id == organization_id)
        if key in seen or existing.first():
            duplicates += 1
            continue
        seen.add(key)
        db.add(Holiday(
            name=name, date=h_date, type=HOLIDAY_TYPES[type_key],
            is_recurring=_truthy(safe.get("is_recurring", safe.get("recurring", False))),
            description=(str(safe.get("description")).strip() or None) if safe.get("description") is not None else None,
            created_by=created_by, organization_id=organization_id,
        ))
        imported += 1
    db.commit()
    return {"total": len(holidays), "imported": imported, "duplicates": duplicates, "errors": errors}


# ── ANALYTICS ────────────────────────────────────────────────────────────────

def _analytics_range(date_from: Optional[date], date_to: Optional[date], default_days: int):
    date_to = date_to or date.today()
    date_from = date_from or (date_to - timedelta(days=default_days))
    return date_from, date_to


def _range_records(db: Session, date_from: date, date_to: date, organization_id: Optional[int]):
    """One record per person per day (a duplicated record must not double anything), this
    organization only, soft-deleted rows excluded."""
    query = db.query(AttendanceRecord).filter(
        AttendanceRecord.date >= date_from, AttendanceRecord.date <= date_to, AttendanceRecord.is_deleted.isnot(True),
    )
    if organization_id is not None:
        query = query.filter(AttendanceRecord.organization_id == organization_id)
    unique: dict = {}
    for r in query.order_by(AttendanceRecord.id).all():
        unique[(r.employee_id, r.date)] = r
    return list(unique.values())


def _overtime_hours(rec) -> float:
    """Hours past the standard day, from the real check-in/out (0 when there is no check-out)."""
    if not rec.check_out:
        return 0.0
    hours = _record_hours(rec)
    return max(hours - STANDARD_WORKDAY_HOURS, 0.0) if hours is not None else 0.0


def get_attendance_trends(db: Session, date_from: Optional[date] = None, date_to: Optional[date] = None, organization_id: Optional[int] = None) -> dict:
    date_from, date_to = _analytics_range(date_from, date_to, 30)
    daily: dict = {}
    for r in _range_records(db, date_from, date_to, organization_id):
        row = daily.setdefault(r.date, {"date": r.date.isoformat(), "label": f"{r.date.day} {r.date.strftime('%b')}", "present": 0, "absent": 0})
        if r.status in ATTENDED_STATUSES:
            row["present"] += 1
        elif r.status == AttendanceStatus.ABSENT:
            row["absent"] += 1
    return {"trends": [daily[k] for k in sorted(daily)], "date_from": str(date_from), "date_to": str(date_to)}


def get_department_analysis(db: Session, date_from: Optional[date] = None, date_to: Optional[date] = None, organization_id: Optional[int] = None) -> dict:
    date_from, date_to = _analytics_range(date_from, date_to, 30)
    records = _range_records(db, date_from, date_to, organization_id)
    # only the people who appear in these records, and only their departments: no whole-table scans
    people = {r.employee_id for r in records}
    dept_ids = {}
    if people:
        dept_ids = dict(db.query(Employee.id, Employee.department_id).filter(Employee.id.in_(people)).all())
    wanted = {d for d in dept_ids.values() if d}
    dept_names = {d.id: d.name for d in db.query(Department).filter(Department.id.in_(wanted)).all()} if wanted else {}
    dept_of = {emp_id: dept_names.get(dept_id) for emp_id, dept_id in dept_ids.items()}
    data: dict = {}
    for r in records:
        if r.status in (AttendanceStatus.ON_LEAVE, AttendanceStatus.HOLIDAY):
            continue  # not a working day for attendance purposes
        row = data.setdefault(dept_of.get(r.employee_id) or "Unassigned", {"present": 0, "absent": 0, "late": 0, "total": 0})
        row["total"] += 1
        if r.status in ATTENDED_STATUSES:
            row["present"] += 1
        if r.status == AttendanceStatus.ABSENT:
            row["absent"] += 1
        if r.status == AttendanceStatus.LATE:
            row["late"] += 1
    breakdown = [
        {"department": d, "present": v["present"], "absent": v["absent"], "late": v["late"], "total_records": v["total"],
         "attendance_rate": round(v["present"] / v["total"] * 100, 2) if v["total"] else 0.0}
        for d, v in sorted(data.items())
    ]
    return {"department_breakdown": breakdown, "total_departments": len(breakdown)}


def _months_between(date_from: date, date_to: date):
    y, m = date_from.year, date_from.month
    while (y, m) <= (date_to.year, date_to.month):
        yield y, m
        y, m = (y + 1, 1) if m == 12 else (y, m + 1)


def get_overtime_analytics(db: Session, date_from: Optional[date] = None, date_to: Optional[date] = None, organization_id: Optional[int] = None) -> dict:
    """Overtime hours per month over the range (default: the last six months). Months with no
    overtime are returned as 0, never invented; hours come from real check-in/out times."""
    if not date_from and not date_to:
        today = date.today()
        date_to = today
        y, m = today.year, today.month - 5
        while m < 1:
            y, m = y - 1, m + 12
        date_from = date(y, m, 1)
    date_from, date_to = _analytics_range(date_from, date_to, 180)
    monthly = {
        (y, m): {"month": f"{y}-{m:02d}", "label": date(y, m, 1).strftime("%b %Y"), "hours": 0.0, "employees": set(), "record_count": 0}
        for y, m in _months_between(date_from, date_to)
    }
    for r in _range_records(db, date_from, date_to, organization_id):
        bucket = monthly.get((r.date.year, r.date.month))
        extra = _overtime_hours(r)
        if bucket is not None and extra > 0:
            bucket["hours"] += extra
            bucket["employees"].add(r.employee_id)
            bucket["record_count"] += 1
    rows = []
    for b in monthly.values():
        hours = round(b["hours"], 2)
        rows.append({"month": b["month"], "label": b["label"], "hours": hours, "total_hours": hours,
                     "employees": len(b["employees"]), "record_count": b["record_count"]})
    return {"monthly_breakdown": rows, "total_months": len(rows), "total_hours": round(sum(r["hours"] for r in rows), 2),
            "date_from": str(date_from), "date_to": str(date_to)}


def _roster_outcomes(db: Session, date_from: date, date_to: date, organization_id: Optional[int]):
    """(shift name, was the rostered person in that day) for each roster row in range."""
    query = db.query(ShiftRoster, Shift).join(Shift, ShiftRoster.shift_id == Shift.id).filter(
        ShiftRoster.date >= date_from, ShiftRoster.date <= date_to, ShiftRoster.is_active.is_(True),
    )
    if organization_id is not None:
        query = query.filter(Shift.organization_id == organization_id)
    attended = {(r.employee_id, r.date) for r in _range_records(db, date_from, date_to, organization_id) if r.status in ATTENDED_STATUSES}
    return [(shift.name or "Unknown", (roster.employee_id, roster.date) in attended) for roster, shift in query.all()]


def get_shift_efficiency(db: Session, date_from: Optional[date] = None, date_to: Optional[date] = None, organization_id: Optional[int] = None) -> dict:
    date_from, date_to = _analytics_range(date_from, date_to, 30)
    data: dict = {}
    for name, was_in in _roster_outcomes(db, date_from, date_to, organization_id):
        row = data.setdefault(name, {"total_assigned": 0, "total_present": 0})
        row["total_assigned"] += 1
        row["total_present"] += 1 if was_in else 0
    efficiency = [
        {"shift": s, "total_assigned": v["total_assigned"], "total_present": v["total_present"],
         "efficiency": round(v["total_present"] / v["total_assigned"] * 100, 2) if v["total_assigned"] else 0.0}
        for s, v in sorted(data.items())
    ]
    return {"shift_efficiency": efficiency, "total_shifts": len(efficiency)}


def _period_kpis(db: Session, date_from: date, date_to: date, organization_id: Optional[int]) -> dict:
    records = _range_records(db, date_from, date_to, organization_id)
    attended = sum(1 for r in records if r.status in ATTENDED_STATUSES)
    absent = sum(1 for r in records if r.status == AttendanceStatus.ABSENT)
    hours = [h for h in (_record_hours(r) for r in records if r.check_out) if h is not None]
    outcomes = _roster_outcomes(db, date_from, date_to, organization_id)
    return {
        "attendance_rate": round(attended / (attended + absent) * 100, 1) if attended + absent else None,
        "avg_work_hours": round(sum(hours) / len(hours), 2) if hours else None,
        "total_overtime": round(sum(_overtime_hours(r) for r in records), 2),
        "shift_efficiency": round(sum(1 for _, ok in outcomes if ok) / len(outcomes) * 100, 1) if outcomes else None,
        "records": len(records),
    }


def get_attendance_kpis(db: Session, date_from: Optional[date] = None, date_to: Optional[date] = None, organization_id: Optional[int] = None) -> dict:
    """The four headline figures for the range, each compared with the equally long period before it.
    A figure with nothing behind it is None (shown as a dash), never a placeholder number."""
    date_from, date_to = _analytics_range(date_from, date_to, 30)
    span = (date_to - date_from).days + 1
    prev_to = date_from - timedelta(days=1)
    prev_from = prev_to - timedelta(days=span - 1)
    now = _period_kpis(db, date_from, date_to, organization_id)
    before = _period_kpis(db, prev_from, prev_to, organization_id)
    changes = {k: _change(now[k], before[k]) if now[k] is not None and before[k] else None
               for k in ("attendance_rate", "avg_work_hours", "total_overtime", "shift_efficiency")}
    return {**now, "changes": changes, "date_from": str(date_from), "date_to": str(date_to)}


# ── ATTENDANCE ANALYTICS ───────────────────────────────────────────────────────────────

def get_attendance_analytics(db: Session, date_from: Optional[date] = None, date_to: Optional[date] = None, organization_id: Optional[int] = None) -> dict:
    today = date.today()
    org_filter = [AttendanceRecord.organization_id == organization_id] if organization_id else []
    
    present_today = db.query(func.count(AttendanceRecord.id)).filter(
        AttendanceRecord.date == today,
        AttendanceRecord.status == AttendanceStatus.PRESENT,
        *org_filter,
    ).scalar() or 0
    
    absent_today = db.query(func.count(AttendanceRecord.id)).filter(
        AttendanceRecord.date == today,
        AttendanceRecord.status == AttendanceStatus.ABSENT,
        *org_filter,
    ).scalar() or 0
    
    on_leave_count = db.query(func.count(AttendanceRecord.id)).filter(
        AttendanceRecord.date == today,
        AttendanceRecord.status == AttendanceStatus.ON_LEAVE,
        *org_filter,
    ).scalar() or 0
    
    remote_count = db.query(func.count(AttendanceRecord.id)).filter(
        AttendanceRecord.date == today,
        AttendanceRecord.status == AttendanceStatus.REMOTE,
        *org_filter,
    ).scalar() or 0
    
    late_arrivals = db.query(func.count(AttendanceRecord.id)).filter(
        AttendanceRecord.date == today,
        AttendanceRecord.status == AttendanceStatus.LATE,
        *org_filter,
    ).scalar() or 0
    
    emp_filter = [Employee.organization_id == organization_id] if organization_id else []
    total_emp = db.query(func.count(Employee.id)).filter(
        Employee.is_active == True, *emp_filter
    ).scalar() or 1
    attendance_percentage = round((present_today / total_emp) * 100, 2) if total_emp else 0.0
    
    avg_hours = db.query(
        func.avg(
            func.extract("epoch", AttendanceRecord.check_out - AttendanceRecord.check_in) / 3600
        )
    ).filter(
        AttendanceRecord.date == today,
        AttendanceRecord.check_in.isnot(None),
        AttendanceRecord.check_out.isnot(None),
        *org_filter,
    ).scalar() or 0.0
    avg_working_hours = round(float(avg_hours), 2)
    
    return {
        "present_today": present_today,
        "absent_today": absent_today,
        "late_arrivals": late_arrivals,
        "early_departures": 0,
        "on_leave_count": on_leave_count,
        "remote_count": remote_count,
        "overtime_count": 0,
        "attendance_percentage": attendance_percentage,
        "avg_working_hours": avg_working_hours,
        "total_employees": total_emp,
    }


# ── LEAVE MANAGEMENT ──────────────────────────────────────────────────────────────────

def get_leave_dashboard(db: Session, organization_id: Optional[int] = None) -> dict:
    today = date.today()
    base_filter = [LeaveRequest.organization_id == organization_id] if organization_id else []
    total_requests = db.query(func.count(LeaveRequest.id)).filter(*base_filter).scalar() or 0
    pending_requests = db.query(func.count(LeaveRequest.id)).filter(
        LeaveRequest.status == RequestStatus.PENDING, *base_filter
    ).scalar() or 0
    approved_requests = db.query(func.count(LeaveRequest.id)).filter(
        LeaveRequest.status == RequestStatus.APPROVED, *base_filter
    ).scalar() or 0
    rejected_requests = db.query(func.count(LeaveRequest.id)).filter(
        LeaveRequest.status == RequestStatus.REJECTED, *base_filter
    ).scalar() or 0
    total_days_taken = db.query(func.coalesce(func.sum(LeaveRequest.total_days), 0)).filter(
        LeaveRequest.status == RequestStatus.APPROVED, *base_filter
    ).scalar() or 0
    on_leave_today = db.query(func.count(LeaveRequest.id)).filter(
        LeaveRequest.start_date <= today,
        LeaveRequest.end_date >= today,
        LeaveRequest.status == RequestStatus.APPROVED,
        *base_filter,
    ).scalar() or 0
    wfh = db.query(func.count(LeaveRequest.id)).filter(
        LeaveRequest.start_date <= today,
        LeaveRequest.end_date >= today,
        LeaveRequest.status == RequestStatus.APPROVED,
        LeaveRequest.leave_type == LeaveType.WORK_FROM_HOME,
        *base_filter,
    ).scalar() or 0
    emp_filter = [Employee.organization_id == organization_id] if organization_id else []
    employee_count = db.query(func.count(Employee.id)).filter(
        Employee.is_deleted == False, *emp_filter
    ).scalar() or 0
    return {
        "employee_count": employee_count,
        "total_requests": total_requests,
        "pending_requests": pending_requests,
        "approved_requests": approved_requests,
        "rejected_requests": rejected_requests,
        "total_days_taken": total_days_taken,
        "on_leave_today": on_leave_today,
        "wfh": wfh,
    }


def get_leave_requests(
    db: Session,
    page: int = 1,
    per_page: int = 20,
    employee_id: Optional[int] = None,
    status: Optional[str] = None,
    leave_type: Optional[str] = None,
    organization_id: Optional[int] = None,
) -> dict:
    per_page = min(per_page, 100)
    query = db.query(LeaveRequest)
    
    if organization_id:
        query = query.filter(LeaveRequest.organization_id == organization_id)
    if employee_id:
        query = query.filter(LeaveRequest.employee_id == employee_id)
    if status:
        query = query.filter(LeaveRequest.status == status)
    if leave_type:
        query = query.filter(LeaveRequest.leave_type == leave_type)
    
    total = query.count()
    items = query.options(selectinload(LeaveRequest.employee)).order_by(LeaveRequest.created_at.desc()).offset((page - 1) * per_page).limit(per_page).all()
    
    result = []
    for r in items:
        result.append({
            "id": r.id,
            "employee_id": r.employee_id,
            "leave_type": r.leave_type,
            "start_date": r.start_date,
            "end_date": r.end_date,
            "days": r.days,
            "reason": r.reason,
            "status": r.status,
            "reviewed_by": r.reviewed_by,
            "reviewed_at": r.reviewed_at,
            "created_at": r.created_at,
            "updated_at": r.updated_at,
            "employee_name": r.employee.full_name if r.employee else None,
        })
    
    return {"total": total, "page": page, "per_page": per_page, "items": result}


def get_leave_request_by_id(db: Session, leave_id: int, organization_id: Optional[int] = None) -> LeaveRequest:
    leave = db.query(LeaveRequest).filter(LeaveRequest.id == leave_id).first()
    if not leave:
        raise NotFoundException("LeaveRequest", leave_id)
    if organization_id and leave.organization_id != organization_id:
        raise NotFoundException("LeaveRequest", leave_id)
    return leave


def _leave_activity(db: Session, action_type: str, actor, leave: LeaveRequest, changes=None, target=None) -> None:
    """Same activity events as the main leave service, so the Super Admin feed is
    identical whichever leave endpoint the organization used."""
    from app.modules.hr import service as hr_service

    hr_service._activity_record(
        db, action_type, hr_service._actor_of(db, actor), leave.organization_id,
        entity_type="LeaveRequest", entity_id=leave.id, **(target or hr_service._leave_target(db, leave)),
        changes=changes,
    )


def create_leave_request(db: Session, data: dict, created_by: int = None, organization_id: Optional[int] = None) -> LeaveRequest:
    if organization_id:
        data["organization_id"] = organization_id
    leave = LeaveRequest(**data)  # LeaveRequest has no created_by column; created_by is the activity actor
    db.add(leave)
    db.commit()
    db.refresh(leave)
    _leave_activity(db, "leave.requested", created_by, leave, changes=[
        {"field": "days", "label": "Days", "before": None, "after": leave.days},
        {"field": "status", "label": "Status", "before": None, "after": getattr(leave.status, "value", leave.status)},
    ])
    return leave


def update_leave_request(db: Session, leave_id: int, data: dict, reviewed_by: int = None, organization_id: Optional[int] = None) -> LeaveRequest:
    from app.modules.super_admin.activity_service import changes_from, snapshot

    leave = get_leave_request_by_id(db, leave_id, organization_id)
    update_data = sanitize_dict(data)
    before = snapshot(leave, update_data.keys())
    for field, value in update_data.items():
        setattr(leave, field, value)
    if reviewed_by:
        leave.reviewed_by = reviewed_by
        leave.reviewed_at = func.now()
    db.commit()
    db.refresh(leave)
    _leave_activity(db, "leave.updated", reviewed_by, leave, changes=changes_from(before, update_data))
    return leave


def delete_leave_request(db: Session, leave_id: int, organization_id: Optional[int] = None, actor=None) -> None:
    from app.modules.hr import service as hr_service

    leave = get_leave_request_by_id(db, leave_id, organization_id)
    target = hr_service._leave_target(db, leave)  # read before the row is gone
    gone = LeaveRequest(id=leave.id, organization_id=leave.organization_id, employee_id=leave.employee_id)
    db.delete(leave)
    db.commit()
    _leave_activity(db, "leave.deleted", actor, gone, target=target)


def review_leave_request(db: Session, leave_id: int, data: dict, reviewed_by: int = None, organization_id: Optional[int] = None) -> LeaveRequest:
    leave = get_leave_request_by_id(db, leave_id, organization_id)
    before_status = getattr(leave.status, "value", leave.status)
    if reviewed_by:
        leave.reviewed_by = reviewed_by
        leave.reviewed_at = func.now()
    update_data = sanitize_dict(data)
    for field, value in update_data.items():
        setattr(leave, field, value)
    db.commit()
    db.refresh(leave)
    if leave.status in (RequestStatus.APPROVED, RequestStatus.REJECTED):
        _leave_activity(
            db, "leave.approved" if leave.status == RequestStatus.APPROVED else "leave.rejected", reviewed_by, leave,
            changes=[{"field": "status", "label": "Status", "before": before_status,
                      "after": getattr(leave.status, "value", leave.status)}],
        )

    # Notify the employee about the review decision (non-blocking)
    if leave.status in (RequestStatus.APPROVED, RequestStatus.REJECTED):
        try:
            from app.services.email_service import send_leave_approved, send_leave_rejected
            from app.config import settings
            employee = db.query(Employee).filter(Employee.id == leave.employee_id).first()
            org = db.query(Organization).filter(Organization.id == leave.organization_id).first()
            if employee:
                if leave.status == RequestStatus.APPROVED:
                    send_leave_approved(
                        email=employee.email,
                        first_name=employee.first_name or "there",
                        request_reference=f"LV-{leave.organization_id}-{leave.id:04d}",
                        leave_period_display=(
                            f"{leave.start_date.strftime('%b %d, %Y')} – {leave.end_date.strftime('%b %d, %Y')}"
                        ),
                        leave_request_url=f"{settings.FRONTEND_URL.rstrip('/')}/employee/leaves",
                        workspace_name=(org.organization_name or org.display_name or "Your workspace") if org else "Your workspace",
                        db=db,
                        organization_id=leave.organization_id,
                    )
                else:
                    send_leave_rejected(
                        email=employee.email,
                        first_name=employee.first_name or "there",
                        request_reference=f"LV-{leave.organization_id}-{leave.id:04d}",
                        leave_request_url=f"{settings.FRONTEND_URL.rstrip('/')}/employee/leaves",
                        db=db,
                        organization_id=leave.organization_id,
                    )
        except Exception as e:
            logger.warning(f"[email] Failed to notify employee about reviewed leave request {leave.id}: {e}")

    return leave


def get_leave_balance(db: Session, employee_id: Optional[int] = None, organization_id: Optional[int] = None) -> dict:
    query = db.query(LeaveBalance)
    if organization_id:
        query = query.filter(LeaveBalance.organization_id == organization_id)
    if employee_id:
        query = query.filter(LeaveBalance.employee_id == employee_id)
    
    balances = query.all()
    result = {}
    for b in balances:
        # b.leave_type is a str-mixin Enum; str() on it renders "LeaveType.SICK"
        # rather than "sick" on this Python version, so use .value explicitly.
        leave_type_key = b.leave_type.value if hasattr(b.leave_type, "value") else str(b.leave_type)
        result[leave_type_key] = {
            "total_days": b.total_days,
            "used_days": b.used_days,
            "pending_days": b.pending_days,
            "year": b.year,
            "employee_id": b.employee_id,
            "organization_id": b.organization_id,
        }
    
    return result


def init_leave_balance(db: Session, employee_id: int, year: int, created_by: int = None, organization_id: Optional[int] = None) -> dict:
    employee = db.query(Employee).filter(Employee.id == employee_id).first()
    if not employee:
        raise NotFoundException("Employee", employee_id)
    if organization_id and employee.organization_id != organization_id:
        raise NotFoundException("Employee", employee_id)
    
    org_id = organization_id or employee.organization_id or 1
    
    existing = db.query(LeaveBalance).filter(
        LeaveBalance.employee_id == employee_id,
        LeaveBalance.organization_id == org_id,
        LeaveBalance.year == year
    ).first()
    
    if existing:
        return {"message": f"Leave balance already exists for employee {employee_id} in year {year}"}
    
    default_balances = [
        {"leave_type": LeaveType.ANNUAL, "total_days": 21},
        {"leave_type": LeaveType.SICK, "total_days": 12},
        {"leave_type": LeaveType.CASUAL, "total_days": 5},
        {"leave_type": LeaveType.UNPAID, "total_days": 0},
    ]
    
    created = []
    for bal in default_balances:
        balance = LeaveBalance(
            employee_id=employee_id,
            organization_id=org_id,
            leave_type=bal["leave_type"],
            total_days=bal["total_days"],
            used_days=0,
            pending_days=0,
            year=year,
            created_by=created_by,
        )
        db.add(balance)
        created.append(balance)
    
    db.commit()
    for bal in created:
        db.refresh(bal)
    
    return {"created": len(created), "balances": created}


# ── ATTENDANCE EXPORTS ─────────────────────────────────────────────────────────

EXPORT_COLUMNS = ["Employee Code", "Employee Name", "Department", "Date", "Check In", "Check Out", "Status", "Hours Worked", "Notes"]


def _hours_worked(rec) -> Optional[float]:
    """The stored total, or check-out minus check-in when the record has no stored total."""
    if rec.total_hours is not None:
        return float(rec.total_hours)
    if rec.check_in and rec.check_out and rec.check_out >= rec.check_in:
        return round((rec.check_out - rec.check_in).total_seconds() / 3600, 2)
    return None


def _export_rows(db: Session, organization_id: Optional[int], date_from, date_to, employee_id):
    """Attendance rows for an export: this organization only, soft-deleted rows excluded, with
    the employee's code, name and department instead of a bare database id."""
    query = (
        db.query(AttendanceRecord, Employee, Department)
        .join(Employee, Employee.id == AttendanceRecord.employee_id)
        .outerjoin(Department, Department.id == Employee.department_id)
        .filter(AttendanceRecord.is_deleted.isnot(True))
    )
    if organization_id is not None:
        query = query.filter(AttendanceRecord.organization_id == organization_id)
    if date_from:
        query = query.filter(AttendanceRecord.date >= date_from)
    if date_to:
        query = query.filter(AttendanceRecord.date <= date_to)
    if employee_id:
        query = query.filter(AttendanceRecord.employee_id == employee_id)
    rows = []
    for rec, emp, dept in query.order_by(AttendanceRecord.date.desc(), Employee.first_name, Employee.last_name).all():
        rows.append({
            "code": emp.employee_code or "",
            "name": f"{emp.first_name or ''} {emp.last_name or ''}".strip(),
            "department": (dept.name if dept else "") or "",
            "date": rec.date,
            "check_in": rec.check_in,
            "check_out": rec.check_out,
            "status": rec.status.value if rec.status else "",
            "hours": _hours_worked(rec),
            "notes": rec.notes or "",
        })
    return rows


def _csv_safe(value):
    """A spreadsheet runs a cell that starts with = + - @ as a formula; neutralise names and notes."""
    text = "" if value is None else str(value)
    return "'" + text if text[:1] in ("=", "+", "-", "@", "\t", "\r") else text


def _export_filename(date_from, date_to, ext: str) -> str:
    return f"attendance_{date_from or 'all'}_{date_to or 'all'}.{ext}"


def export_attendance_csv(
    db: Session,
    date_from: Optional[date] = None,
    date_to: Optional[date] = None,
    employee_id: Optional[int] = None,
    organization_id: Optional[int] = None,
):
    """A real CSV: plain text, one line per record, ISO dates and 24h HH:MM times. The UTF-8 BOM
    lets Excel open names with accents correctly."""
    import csv
    import io

    rows = _export_rows(db, organization_id, date_from, date_to, employee_id)
    output = io.StringIO()
    writer = csv.writer(output, lineterminator="\r\n")
    writer.writerow(EXPORT_COLUMNS)
    for r in rows:
        writer.writerow([
            _csv_safe(r["code"]), _csv_safe(r["name"]), _csv_safe(r["department"]),
            r["date"].isoformat() if r["date"] else "",
            r["check_in"].strftime("%H:%M") if r["check_in"] else "",
            r["check_out"].strftime("%H:%M") if r["check_out"] else "",
            r["status"], "" if r["hours"] is None else f"{r['hours']:.2f}", _csv_safe(r["notes"]),
        ])
    from fastapi.responses import Response
    return Response(
        content=("\ufeff" + output.getvalue()).encode("utf-8"),
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": f'attachment; filename="{_export_filename(date_from, date_to, "csv")}"'},
    )


def export_attendance_excel(
    db: Session,
    date_from: Optional[date] = None,
    date_to: Optional[date] = None,
    employee_id: Optional[int] = None,
    organization_id: Optional[int] = None,
):
    """A real workbook: typed date / time / number cells, a styled frozen header with filters, sized
    columns, and a Summary sheet with totals by status."""
    from io import BytesIO

    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    rows = _export_rows(db, organization_id, date_from, date_to, employee_id)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Attendance"
    ws.append(EXPORT_COLUMNS)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F3A5F")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for r in rows:
        ws.append([
            r["code"], r["name"], r["department"], r["date"],
            r["check_in"].time() if r["check_in"] else None,
            r["check_out"].time() if r["check_out"] else None,
            r["status"].replace("_", " ").title(), r["hours"], r["notes"],
        ])
    for row in ws.iter_rows(min_row=2):
        row[3].number_format = "DD-MMM-YYYY"
        row[4].number_format = "HH:MM"
        row[5].number_format = "HH:MM"
        row[7].number_format = "0.00"
    for idx, col in enumerate(EXPORT_COLUMNS, start=1):
        longest = max([len(str(col))] + [len(str(c.value)) for c in ws[get_column_letter(idx)][1:200] if c.value is not None])
        ws.column_dimensions[get_column_letter(idx)].width = min(max(longest + 2, 10), 40)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    summary = wb.create_sheet("Summary")
    summary.append(["Status", "Records"])
    for cell in summary[1]:
        cell.font = Font(bold=True)
    counts = {}
    for r in rows:
        key = r["status"].replace("_", " ").title() or "Unknown"
        counts[key] = counts.get(key, 0) + 1
    for key in sorted(counts):
        summary.append([key, counts[key]])
    summary.append(["Total", len(rows)])
    summary.cell(row=summary.max_row, column=1).font = Font(bold=True)
    summary.cell(row=summary.max_row, column=2).font = Font(bold=True)
    summary.column_dimensions["A"].width = 22
    summary.column_dimensions["B"].width = 12

    output = BytesIO()
    wb.save(output)
    from fastapi.responses import Response
    return Response(
        content=output.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{_export_filename(date_from, date_to, "xlsx")}"'},
    )
