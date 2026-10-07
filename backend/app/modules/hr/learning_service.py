from datetime import datetime, date
from typing import Optional
import json
from sqlalchemy import func
from sqlalchemy.orm import Session

from app.modules.hr.models import (
    LearningCourse, LearningEnrollment, LearningPath, LearningPathItem,
    LearningCertification, LearningSkill, LearningAssessment,
    LearningAssessmentQuestion, LearningQuizAttempt,
    LearningTrainingProgram, LearningTrainingProgramAssignment,
    LearningCalendarEvent,
)
from app.modules.hr.schemas import (
    CourseCreate, CourseUpdate,
    EnrollmentCreate, EnrollmentUpdate,
    LearningPathCreate, LearningPathUpdate, LearningPathItemCreate, LearningPathItemResponse,
    CertificationCreate, CertificationUpdate,
    SkillCreate, SkillUpdate,
    AssessmentCreate, AssessmentUpdate,
    QuestionCreate, QuestionUpdate,
    QuizAttemptStart, QuizAttemptSubmit,
    TrainingProgramCreate, TrainingProgramUpdate,
    ProgramAssignmentCreate, ProgramAssignmentUpdate,
    CalendarEventCreate, CalendarEventUpdate,
)
from sqlalchemy.exc import SQLAlchemyError
from app.core.exceptions import NotFoundException, BadRequestException


def _check_course_name(db: Session, organization_id, name: str, exclude_id=None) -> None:
    q = db.query(LearningCourse.id).filter(func.lower(LearningCourse.course_name) == name.lower())
    if organization_id is not None:
        q = q.filter(LearningCourse.organization_id == organization_id)
    if exclude_id:
        q = q.filter(LearningCourse.id != exclude_id)
    if q.first():
        raise BadRequestException(f"A course named '{name}' already exists. Use a different name or edit that course.")


def create_course(db: Session, data: CourseCreate, created_by: int = None, organization_id: Optional[int] = None) -> LearningCourse:
    _check_course_name(db, organization_id, data.course_name)
    course = LearningCourse(**data.model_dump(), created_by=created_by)
    if organization_id is not None:
        course.organization_id = organization_id
    db.add(course)
    db.commit()
    db.refresh(course)
    return course


def get_courses(
    db: Session,
    page: int = 1,
    per_page: int = 20,
    search: Optional[str] = None,
    category: Optional[str] = None,
    status: Optional[str] = None,
    course_type: Optional[str] = None,
    organization_id: Optional[int] = None,
) -> dict:
    per_page = min(per_page, 100)
    query = db.query(LearningCourse)

    if search:
        term = f"%{search}%"
        query = query.filter(LearningCourse.course_name.ilike(term))

    if category:
        query = query.filter(LearningCourse.category == category)

    if status:
        query = query.filter(LearningCourse.status == status)

    if course_type:
        query = query.filter(LearningCourse.course_type == course_type)

    if organization_id is not None:
        query = query.filter(LearningCourse.organization_id == organization_id)

    total = query.count()
    courses = query.order_by(LearningCourse.created_at.desc()).offset((page - 1) * per_page).limit(per_page).all()

    return {
        "total": total,
        "page": page,
        "per_page": per_page,
        "items": courses,
    }


def get_course_by_id(db: Session, course_id: int, organization_id: Optional[int] = None) -> LearningCourse:
    query = db.query(LearningCourse).filter(LearningCourse.id == course_id)
    if organization_id is not None:
        query = query.filter(LearningCourse.organization_id == organization_id)   # another organization's course does not exist here
    course = query.first()
    if not course:
        raise NotFoundException("LearningCourse", course_id)
    return course


def update_course(db: Session, course_id: int, data: CourseUpdate, organization_id: Optional[int] = None) -> LearningCourse:
    course = get_course_by_id(db, course_id, organization_id)
    update_data = data.model_dump(exclude_unset=True)
    if update_data.get("course_name") and update_data["course_name"].lower() != (course.course_name or "").lower():
        _check_course_name(db, organization_id, update_data["course_name"], exclude_id=course.id)
    for field, value in update_data.items():
        setattr(course, field, value)
    db.commit()
    db.refresh(course)
    return course


def delete_course(db: Session, course_id: int, organization_id: Optional[int] = None) -> None:
    course = get_course_by_id(db, course_id, organization_id)
    enrolled = db.query(LearningEnrollment.id).filter(LearningEnrollment.course_id == course.id).count()
    if enrolled:
        raise BadRequestException(
            f"This course has {enrolled} enrollment{'s' if enrolled != 1 else ''} and cannot be deleted. Set it to Inactive instead to hide it from learners."
        )
    db.delete(course)
    db.commit()


def create_enrollment(db: Session, data: EnrollmentCreate, organization_id: Optional[int] = None) -> LearningEnrollment:
    get_course_by_id(db, data.course_id)
    enrollment = LearningEnrollment(**data.model_dump())
    if organization_id is not None:
        enrollment.organization_id = organization_id
    db.add(enrollment)
    db.commit()
    db.refresh(enrollment)
    return enrollment


def get_enrollments(
    db: Session,
    page: int = 1,
    per_page: int = 20,
    employee_id: Optional[int] = None,
    course_id: Optional[int] = None,
    status: Optional[str] = None,
    organization_id: Optional[int] = None,
) -> dict:
    per_page = min(per_page, 100)
    query = db.query(LearningEnrollment)

    if employee_id:
        query = query.filter(LearningEnrollment.employee_id == employee_id)

    if course_id:
        query = query.filter(LearningEnrollment.course_id == course_id)

    if status:
        query = query.filter(LearningEnrollment.status == status)

    if organization_id is not None:
        query = query.filter(LearningEnrollment.organization_id == organization_id)

    total = query.count()
    enrollments = query.order_by(LearningEnrollment.enrolled_at.desc()).offset((page - 1) * per_page).limit(per_page).all()

    items = []
    for e in enrollments:
        items.append({
            "id": e.id,
            "course_id": e.course_id,
            "employee_id": e.employee_id,
            "status": e.status,
            "progress_pct": e.progress_pct,
            "enrolled_at": e.enrolled_at,
            "started_at": e.started_at,
            "completed_at": e.completed_at,
            "score": e.score,
            "notes": e.notes,
            "created_at": e.created_at,
            "updated_at": e.updated_at,
            "course_name": e.course.course_name if e.course else None,
            "employee_name": e.employee.full_name if e.employee else None,
        })

    return {
        "total": total,
        "page": page,
        "per_page": per_page,
        "items": items,
    }


def get_enrollment_by_id(db: Session, enrollment_id: int) -> LearningEnrollment:
    enrollment = db.query(LearningEnrollment).filter(LearningEnrollment.id == enrollment_id).first()
    if not enrollment:
        raise NotFoundException("LearningEnrollment", enrollment_id)
    return enrollment


def update_enrollment(db: Session, enrollment_id: int, data: EnrollmentUpdate, organization_id: Optional[int] = None) -> LearningEnrollment:
    enrollment = get_enrollment_by_id(db, enrollment_id)
    update_data = data.model_dump(exclude_unset=True)
    for field, value in update_data.items():
        setattr(enrollment, field, value)
    db.commit()
    db.refresh(enrollment)
    return enrollment


def delete_enrollment(db: Session, enrollment_id: int, organization_id: Optional[int] = None) -> None:
    enrollment = get_enrollment_by_id(db, enrollment_id)
    db.delete(enrollment)
    db.commit()


def create_learning_path(db: Session, data: LearningPathCreate, created_by: int, organization_id: Optional[int] = None) -> LearningPath:
    path = LearningPath(**data.model_dump(), created_by=created_by)
    if organization_id is not None:
        path.organization_id = organization_id
    db.add(path)
    db.commit()
    db.refresh(path)
    return path


def get_learning_paths(db: Session, organization_id: Optional[int] = None) -> list[dict]:
    query = db.query(LearningPath)
    if organization_id is not None:
        query = query.filter(LearningPath.organization_id == organization_id)
    paths = query.order_by(LearningPath.created_at.desc()).all()
    result = []
    for path in paths:
        items = []
        for item in path.items:
            items.append({
                "id": item.id,
                "path_id": item.path_id,
                "course_id": item.course_id,
                "sort_order": item.sort_order,
                "is_required": item.is_required,
                "course_name": item.course.course_name if item.course else None,
            })
        result.append({
            "id": path.id,
            "name": path.name,
            "description": path.description,
            "created_by": path.created_by,
            "is_active": path.is_active,
            "created_at": path.created_at,
            "updated_at": path.updated_at,
            "items": items,
        })
    return result


def get_learning_path_by_id(db: Session, path_id: int, organization_id: Optional[int] = None) -> LearningPath:
    path = db.query(LearningPath).filter(LearningPath.id == path_id).first()
    if not path:
        raise NotFoundException("LearningPath", path_id)
    return path


def update_learning_path(db: Session, path_id: int, data: LearningPathUpdate, organization_id: Optional[int] = None) -> LearningPath:
    path = get_learning_path_by_id(db, path_id, organization_id)
    update_data = data.model_dump(exclude_unset=True)
    for field, value in update_data.items():
        setattr(path, field, value)
    db.commit()
    db.refresh(path)
    return path


def delete_learning_path(db: Session, path_id: int, organization_id: Optional[int] = None) -> None:
    path = get_learning_path_by_id(db, path_id, organization_id)
    db.delete(path)
    db.commit()


def add_path_item(db: Session, path_id: int, data: LearningPathItemCreate, organization_id: Optional[int] = None) -> LearningPathItem:
    get_learning_path_by_id(db, path_id, organization_id)
    get_course_by_id(db, data.course_id)
    item = LearningPathItem(**data.model_dump(), path_id=path_id)
    if organization_id is not None:
        item.organization_id = organization_id
    db.add(item)
    db.commit()
    db.refresh(item)
    return item


def remove_path_item(db: Session, path_id: int, item_id: int, organization_id: Optional[int] = None) -> None:
    get_learning_path_by_id(db, path_id, organization_id)
    item = db.query(LearningPathItem).filter(LearningPathItem.id == item_id, LearningPathItem.path_id == path_id).first()
    if not item:
        raise NotFoundException("LearningPathItem", item_id)
    db.delete(item)
    db.commit()


def update_path_item(db: Session, path_id: int, item_id: int, data: LearningPathItemCreate, organization_id: Optional[int] = None) -> LearningPathItem:
    get_learning_path_by_id(db, path_id, organization_id)
    item = db.query(LearningPathItem).filter(LearningPathItem.id == item_id, LearningPathItem.path_id == path_id).first()
    if not item:
        raise NotFoundException("LearningPathItem", item_id)
    update_data = data.model_dump(exclude_unset=True)
    for field, value in update_data.items():
        setattr(item, field, value)
    db.commit()
    db.refresh(item)
    return item


def create_certification(db: Session, data: CertificationCreate, created_by: int, organization_id: Optional[int] = None) -> LearningCertification:
    cert = LearningCertification(**data.model_dump(exclude={'created_by'}), created_by=created_by)
    if organization_id is not None:
        cert.organization_id = organization_id
    db.add(cert)
    db.commit()
    db.refresh(cert)
    return cert


def get_certifications(db: Session, employee_id: Optional[int] = None, organization_id: Optional[int] = None) -> list[LearningCertification]:
    query = db.query(LearningCertification)
    if employee_id:
        query = query.filter(LearningCertification.employee_id == employee_id)
    if organization_id is not None:
        query = query.filter(LearningCertification.organization_id == organization_id)
    return query.order_by(LearningCertification.issue_date.desc()).all()


def get_certification_by_id(db: Session, cert_id: int) -> LearningCertification:
    cert = db.query(LearningCertification).filter(LearningCertification.id == cert_id).first()
    if not cert:
        raise NotFoundException("LearningCertification", cert_id)
    return cert


def update_certification(db: Session, cert_id: int, data: CertificationUpdate, organization_id: Optional[int] = None) -> LearningCertification:
    cert = get_certification_by_id(db, cert_id)
    update_data = data.model_dump(exclude_unset=True)
    for field, value in update_data.items():
        setattr(cert, field, value)
    db.commit()
    db.refresh(cert)
    return cert


def delete_certification(db: Session, cert_id: int, organization_id: Optional[int] = None) -> None:
    cert = get_certification_by_id(db, cert_id)
    db.delete(cert)
    db.commit()


def create_skill(db: Session, data: SkillCreate, organization_id: Optional[int] = None) -> LearningSkill:
    skill = LearningSkill(**data.model_dump())
    if organization_id is not None:
        skill.organization_id = organization_id
    db.add(skill)
    db.commit()
    db.refresh(skill)
    return skill


def get_skills(db: Session, employee_id: Optional[int] = None) -> list[LearningSkill]:
    query = db.query(LearningSkill)
    if employee_id:
        query = query.filter(LearningSkill.employee_id == employee_id)
    return query.order_by(LearningSkill.skill_name).all()


def get_skill_by_id(db: Session, skill_id: int) -> LearningSkill:
    skill = db.query(LearningSkill).filter(LearningSkill.id == skill_id).first()
    if not skill:
        raise NotFoundException("LearningSkill", skill_id)
    return skill


def update_skill(db: Session, skill_id: int, data: SkillUpdate) -> LearningSkill:
    skill = get_skill_by_id(db, skill_id)
    update_data = data.model_dump(exclude_unset=True)
    for field, value in update_data.items():
        setattr(skill, field, value)
    db.commit()
    db.refresh(skill)
    return skill


def delete_skill(db: Session, skill_id: int) -> None:
    skill = get_skill_by_id(db, skill_id)
    db.delete(skill)
    db.commit()



def _assessment_row(db: Session, assessment_id: int, organization_id: Optional[int] = None) -> LearningAssessment:
    """An assessment belongs to a course, and the course belongs to an organization: another organization's assessment
    does not exist here."""
    query = db.query(LearningAssessment).filter(LearningAssessment.id == assessment_id)
    if organization_id is not None:
        query = query.join(LearningCourse, LearningCourse.id == LearningAssessment.course_id).filter(LearningCourse.organization_id == organization_id)
    assessment = query.first()
    if not assessment:
        raise NotFoundException("LearningAssessment", assessment_id)
    return assessment


def get_assessment_by_id(db: Session, assessment_id: int, organization_id: Optional[int] = None) -> LearningAssessment:
    return _assessment_row(db, assessment_id, organization_id)


def _assessment_dict(db: Session, a: LearningAssessment) -> dict:
    course = db.query(LearningCourse.course_name).filter(LearningCourse.id == a.course_id).first()
    return {
        "id": a.id,
        "course_id": a.course_id,
        "course_name": course[0] if course else None,
        "title": a.title,
        "description": a.description,
        "passing_score": a.passing_score,
        "max_attempts": a.max_attempts,
        "duration_minutes": a.duration_minutes,
        "resource_link": a.resource_link,
        "is_active": bool(a.is_active),
        "questions_count": db.query(LearningAssessmentQuestion).filter(LearningAssessmentQuestion.assessment_id == a.id).count(),
        "attempts_count": db.query(LearningQuizAttempt).filter(LearningQuizAttempt.assessment_id == a.id).count(),
        "created_by": a.created_by,
        "created_at": a.created_at,
        "updated_at": a.updated_at,
    }


def _check_assessment_title(db: Session, course_id: int, title: str, exclude_id=None) -> None:
    q = db.query(LearningAssessment.id).filter(LearningAssessment.course_id == course_id, func.lower(LearningAssessment.title) == title.lower())
    if exclude_id:
        q = q.filter(LearningAssessment.id != exclude_id)
    if q.first():
        raise BadRequestException(f"This course already has an assessment titled '{title}'. Use a different title or edit that assessment.")


def create_assessment(db: Session, data: AssessmentCreate, created_by: int, organization_id: Optional[int] = None) -> dict:
    get_course_by_id(db, data.course_id, organization_id)
    _check_assessment_title(db, data.course_id, data.title)
    assessment = LearningAssessment(**data.model_dump(), created_by=created_by)
    db.add(assessment)
    db.commit()
    db.refresh(assessment)
    return _assessment_dict(db, assessment)


def get_assessments(db: Session, course_id: Optional[int] = None, organization_id: Optional[int] = None, active_only: bool = False) -> list[dict]:
    query = db.query(LearningAssessment)
    if organization_id is not None:
        query = query.join(LearningCourse, LearningCourse.id == LearningAssessment.course_id).filter(LearningCourse.organization_id == organization_id)
    if course_id:
        query = query.filter(LearningAssessment.course_id == course_id)
    if active_only:
        query = query.filter(LearningAssessment.is_active == True)  # noqa: E712
    return [_assessment_dict(db, a) for a in query.order_by(LearningAssessment.created_at.desc()).all()]


def get_assessment_detail(db: Session, assessment_id: int, organization_id: Optional[int] = None) -> dict:
    return _assessment_dict(db, _assessment_row(db, assessment_id, organization_id))


def update_assessment(db: Session, assessment_id: int, data: AssessmentUpdate, organization_id: Optional[int] = None) -> dict:
    assessment = _assessment_row(db, assessment_id, organization_id)
    update_data = data.model_dump(exclude_unset=True)
    for required in ("course_id", "title", "description", "passing_score", "is_active"):
        if required in update_data and update_data[required] is None:
            update_data.pop(required)
    if "course_id" in update_data:
        get_course_by_id(db, update_data["course_id"], organization_id)
    course_id = update_data.get("course_id", assessment.course_id)
    if "title" in update_data or "course_id" in update_data:
        _check_assessment_title(db, course_id, update_data.get("title", assessment.title), exclude_id=assessment.id)
    for field, value in update_data.items():
        setattr(assessment, field, value)
    db.commit()
    db.refresh(assessment)
    return _assessment_dict(db, assessment)


def delete_assessment(db: Session, assessment_id: int, organization_id: Optional[int] = None) -> None:
    assessment = _assessment_row(db, assessment_id, organization_id)
    attempts = db.query(LearningQuizAttempt).filter(LearningQuizAttempt.assessment_id == assessment.id).count()
    if attempts:
        raise BadRequestException(
            f"This assessment has {attempts} recorded attempt{'s' if attempts != 1 else ''} and cannot be deleted. Set it to Inactive instead."
        )
    db.query(LearningAssessmentQuestion).filter(LearningAssessmentQuestion.assessment_id == assessment.id).delete(synchronize_session=False)
    db.delete(assessment)
    db.commit()


def add_question(db: Session, assessment_id: int, data: QuestionCreate, organization_id: Optional[int] = None) -> LearningAssessmentQuestion:
    _assessment_row(db, assessment_id, organization_id)
    values = data.model_dump()
    if values.get("sort_order") is None:
        last = db.query(func.max(LearningAssessmentQuestion.sort_order)).filter(LearningAssessmentQuestion.assessment_id == assessment_id).scalar()
        values["sort_order"] = (last or 0) + 1
    question = LearningAssessmentQuestion(**values, assessment_id=assessment_id)
    db.add(question)
    db.commit()
    db.refresh(question)
    return question


def get_questions(db: Session, assessment_id: int, organization_id: Optional[int] = None) -> list[LearningAssessmentQuestion]:
    if organization_id is not None:
        _assessment_row(db, assessment_id, organization_id)
    return db.query(LearningAssessmentQuestion).filter(
        LearningAssessmentQuestion.assessment_id == assessment_id
    ).order_by(LearningAssessmentQuestion.sort_order, LearningAssessmentQuestion.id).all()


def _question_row(db: Session, assessment_id: int, question_id: int, organization_id: Optional[int]) -> LearningAssessmentQuestion:
    _assessment_row(db, assessment_id, organization_id)
    question = db.query(LearningAssessmentQuestion).filter(
        LearningAssessmentQuestion.id == question_id, LearningAssessmentQuestion.assessment_id == assessment_id
    ).first()
    if not question:
        raise NotFoundException("LearningAssessmentQuestion", question_id)
    return question


def update_question(db: Session, assessment_id: int, question_id: int, data: QuestionUpdate, organization_id: Optional[int] = None) -> LearningAssessmentQuestion:
    from app.modules.hr.schemas import normalize_question
    question = _question_row(db, assessment_id, question_id, organization_id)
    changes = data.model_dump(exclude_unset=True)
    for required in ("question_text", "question_type", "points"):
        if required in changes and changes[required] is None:
            changes.pop(required)
    merged_type = changes.get("question_type", question.question_type)
    raw_options = changes["options"] if "options" in changes else question.options
    raw_answer = changes["correct_answer"] if "correct_answer" in changes else question.correct_answer
    try:
        options, answer = normalize_question(merged_type, raw_options, raw_answer)
    except ValueError as exc:
        raise BadRequestException(str(exc))
    changes.update({"question_type": str(merged_type).strip().lower(), "options": options, "correct_answer": answer})
    for field, value in changes.items():
        setattr(question, field, value)
    db.commit()
    db.refresh(question)
    return question


def delete_question(db: Session, assessment_id: int, question_id: int, organization_id: Optional[int] = None) -> None:
    question = _question_row(db, assessment_id, question_id, organization_id)
    db.delete(question)
    db.commit()


def start_quiz(db: Session, data: QuizAttemptStart, organization_id: Optional[int] = None) -> LearningQuizAttempt:
    assessment = _assessment_row(db, data.assessment_id, organization_id)

    if not assessment.is_active:
        raise BadRequestException("Assessment is not active.")

    if assessment.max_attempts and assessment.max_attempts > 0:
        existing_count = db.query(LearningQuizAttempt).filter(
            LearningQuizAttempt.assessment_id == data.assessment_id,
            LearningQuizAttempt.employee_id == data.employee_id,
        ).count()
        if existing_count >= assessment.max_attempts:
            raise BadRequestException("Maximum number of attempts reached.")

    attempt_number = 1
    last_attempt = db.query(LearningQuizAttempt).filter(
        LearningQuizAttempt.assessment_id == data.assessment_id,
        LearningQuizAttempt.employee_id == data.employee_id,
    ).order_by(LearningQuizAttempt.attempt_number.desc()).first()
    if last_attempt:
        attempt_number = last_attempt.attempt_number + 1

    attempt = LearningQuizAttempt(
        assessment_id=data.assessment_id,
        employee_id=data.employee_id,
        enrollment_id=data.enrollment_id,
        organization_id=organization_id,
        started_at=datetime.utcnow(),
        attempt_number=attempt_number,
        status="in_progress",
    )
    db.add(attempt)
    db.commit()
    db.refresh(attempt)
    return attempt


def submit_quiz(db: Session, attempt_id: int, data: QuizAttemptSubmit, organization_id: Optional[int] = None) -> LearningQuizAttempt:
    attempt = db.query(LearningQuizAttempt).filter(LearningQuizAttempt.id == attempt_id).first()
    if not attempt:
        raise NotFoundException("LearningQuizAttempt", attempt_id)

    if attempt.status == "completed":
        raise BadRequestException("Quiz attempt has already been submitted.")

    assessment = _assessment_row(db, attempt.assessment_id, organization_id)
    questions = get_questions(db, attempt.assessment_id)

    try:
        user_answers = json.loads(data.answers)
    except (json.JSONDecodeError, TypeError):
        raise BadRequestException("Invalid answers format. Expected a JSON array.")

    if not isinstance(user_answers, list):
        raise BadRequestException("Invalid answers format. Expected a JSON array.")

    answer_map = {}
    for entry in user_answers:
        if not isinstance(entry, dict):
            continue
        qid = entry.get("question_id")
        ans = entry.get("answer")
        if qid is not None:
            answer_map[qid] = ans

    total_points = 0
    earned_points = 0

    for question in questions:
        total_points += question.points or 0
        user_answer = answer_map.get(question.id, "")
        if user_answer is not None and question.correct_answer is not None:
            if str(user_answer).strip().lower() == str(question.correct_answer).strip().lower():
                earned_points += question.points or 0

    score = 0
    if total_points > 0:
        score = round((earned_points / total_points) * 100)

    passed = assessment.passing_score is not None and score >= assessment.passing_score

    attempt.answers = data.answers
    attempt.score = score
    attempt.passed = passed
    attempt.completed_at = datetime.utcnow()
    attempt.status = "completed"

    db.commit()
    db.refresh(attempt)
    return attempt


def _attempt_dict(a: LearningQuizAttempt, names: dict) -> dict:
    return {
        "id": a.id, "assessment_id": a.assessment_id, "employee_id": a.employee_id, "employee_name": names.get(a.employee_id),
        "enrollment_id": a.enrollment_id, "started_at": a.started_at, "completed_at": a.completed_at, "score": a.score,
        "passed": a.passed, "answers": a.answers if isinstance(a.answers, str) or a.answers is None else json.dumps(a.answers),
        "attempt_number": a.attempt_number, "status": a.status, "created_at": a.created_at,
    }


def get_quiz_attempts(
    db: Session,
    assessment_id: Optional[int] = None,
    employee_id: Optional[int] = None,
    organization_id: Optional[int] = None,
) -> list[dict]:
    from app.modules.employee.models import Employee
    if assessment_id and organization_id is not None:
        _assessment_row(db, assessment_id, organization_id)
    query = db.query(LearningQuizAttempt)
    if assessment_id:
        query = query.filter(LearningQuizAttempt.assessment_id == assessment_id)
    if employee_id:
        query = query.filter(LearningQuizAttempt.employee_id == employee_id)
    attempts = query.order_by(LearningQuizAttempt.started_at.desc()).all()
    ids = {a.employee_id for a in attempts}
    names = {}
    if ids:
        for e in db.query(Employee).filter(Employee.id.in_(ids)).all():
            names[e.id] = (f"{e.first_name or ''} {e.last_name or ''}".strip()) or e.email
    return [_attempt_dict(a, names) for a in attempts]


def get_quiz_attempt_by_id(db: Session, attempt_id: int, organization_id: Optional[int] = None) -> LearningQuizAttempt:
    attempt = db.query(LearningQuizAttempt).filter(LearningQuizAttempt.id == attempt_id).first()
    if not attempt:
        raise NotFoundException("LearningQuizAttempt", attempt_id)
    if organization_id is not None:
        _assessment_row(db, attempt.assessment_id, organization_id)
    return attempt


# Where a program can go from each status. A finished program stays finished and a cancelled one can only be
# planned again, so it can never flip between completed and cancelled.
PROGRAM_TRANSITIONS = {
    "planned": {"active", "completed", "cancelled"},
    "active": {"completed", "cancelled"},
    "completed": set(),
    "cancelled": {"planned"},
}


def check_program_transition(current: str, new: Optional[str]) -> None:
    if new is None or new == current:
        return
    if new not in PROGRAM_TRANSITIONS.get(current, set()):
        if current == "completed":
            raise BadRequestException("This program is already completed, so its status can no longer be changed.")
        if current == "cancelled":
            raise BadRequestException(f"A cancelled program cannot be marked {new}. Plan it again first.")
        raise BadRequestException(f"A program that is {current} cannot be changed to {new}.")


def _program_row(db: Session, program_id: int, organization_id: Optional[int] = None) -> LearningTrainingProgram:
    query = db.query(LearningTrainingProgram).filter(LearningTrainingProgram.id == program_id)
    if organization_id is not None:
        query = query.filter(LearningTrainingProgram.organization_id == organization_id)   # another organization's program does not exist here
    program = query.first()
    if not program:
        raise NotFoundException("LearningTrainingProgram", program_id)
    return program


def _program_dict(program: LearningTrainingProgram) -> dict:
    instructor = program.instructor
    name = None
    if instructor is not None:
        name = (f"{instructor.first_name or ''} {instructor.last_name or ''}".strip()) or instructor.email
    return {
        "id": program.id,
        "name": program.name,
        "description": program.description,
        "instructor_id": program.instructor_id,
        "instructor_name": name,
        "department": program.department,
        "resource_link": program.resource_link,
        "start_date": program.start_date,
        "end_date": program.end_date,
        "status": program.status,
        "max_participants": program.max_participants,
        "participants_count": len(program.assignments) if program.assignments else 0,
        "created_by": program.created_by,
        "created_at": program.created_at,
        "updated_at": program.updated_at,
    }


def _check_program(db: Session, organization_id, name=None, instructor_id=None, start_date=None, end_date=None, exclude_id=None, check_name=True) -> None:
    if check_name and name:
        q = db.query(LearningTrainingProgram.id).filter(func.lower(LearningTrainingProgram.name) == name.lower())
        if organization_id is not None:
            q = q.filter(LearningTrainingProgram.organization_id == organization_id)
        if exclude_id:
            q = q.filter(LearningTrainingProgram.id != exclude_id)
        if q.first():
            raise BadRequestException(f"A program named '{name}' already exists. Use a different name or edit that program.")
    if instructor_id is not None:
        from app.modules.employee.models import Employee
        q = db.query(Employee.id).filter(Employee.id == instructor_id)
        if organization_id is not None:
            q = q.filter(Employee.organization_id == organization_id)
        if not q.first():
            raise BadRequestException("The selected instructor is not an employee of this organization.")
    if start_date and end_date and end_date < start_date:
        raise BadRequestException("The end date cannot be before the start date.")


def create_training_program(db: Session, data: TrainingProgramCreate, created_by: int, organization_id: Optional[int] = None) -> dict:
    _check_program(db, organization_id, data.name, data.instructor_id, data.start_date, data.end_date)
    program = LearningTrainingProgram(**data.model_dump(), created_by=created_by)
    if organization_id is not None:
        program.organization_id = organization_id
    db.add(program)
    db.commit()
    db.refresh(program)
    return _program_dict(program)


def get_training_programs(
    db: Session,
    page: int = 1,
    per_page: int = 20,
    status: Optional[str] = None,
    organization_id: Optional[int] = None,
    hide_cancelled: bool = False,
) -> dict:
    per_page = min(per_page, 100)
    query = db.query(LearningTrainingProgram)

    if status:
        query = query.filter(LearningTrainingProgram.status == status)
    if hide_cancelled:
        query = query.filter(LearningTrainingProgram.status != "cancelled")
    if organization_id is not None:
        query = query.filter(LearningTrainingProgram.organization_id == organization_id)

    total = query.count()
    programs = query.order_by(LearningTrainingProgram.created_at.desc()).offset((page - 1) * per_page).limit(per_page).all()
    return {"total": total, "page": page, "per_page": per_page, "items": [_program_dict(p) for p in programs]}


def get_training_program_by_id(db: Session, program_id: int, organization_id: Optional[int] = None) -> dict:
    return _program_dict(_program_row(db, program_id, organization_id))


def update_training_program(db: Session, program_id: int, data: TrainingProgramUpdate, organization_id: Optional[int] = None) -> dict:
    program = _program_row(db, program_id, organization_id)
    update_data = data.model_dump(exclude_unset=True)
    check_program_transition(program.status, update_data.get("status"))
    new_name = update_data.get("name")
    _check_program(
        db, organization_id,
        name=new_name,
        instructor_id=update_data.get("instructor_id") if update_data.get("instructor_id") != program.instructor_id else None,
        start_date=update_data.get("start_date", program.start_date),
        end_date=update_data.get("end_date", program.end_date),
        exclude_id=program.id,
        check_name=bool(new_name) and new_name.lower() != (program.name or "").lower(),
    )
    capacity = update_data.get("max_participants")
    if capacity is not None and capacity < (len(program.assignments) if program.assignments else 0):
        raise BadRequestException(f"This program already has {len(program.assignments)} participants, so the maximum cannot be lower than that.")
    for field, value in update_data.items():
        setattr(program, field, value)
    db.commit()
    db.refresh(program)
    return _program_dict(program)


def delete_training_program(db: Session, program_id: int, organization_id: Optional[int] = None) -> None:
    program = _program_row(db, program_id, organization_id)
    if program.assignments:
        n = len(program.assignments)
        raise BadRequestException(
            f"This program has {n} participant{'s' if n != 1 else ''} and cannot be deleted. Set it to Cancelled instead."
        )
    db.delete(program)
    db.commit()


def assign_program(db: Session, data: ProgramAssignmentCreate, organization_id: Optional[int] = None) -> LearningTrainingProgramAssignment:
    get_training_program_by_id(db, data.program_id, organization_id)
    assignment = LearningTrainingProgramAssignment(**data.model_dump())
    if organization_id is not None:
        assignment.organization_id = organization_id
    db.add(assignment)
    db.commit()
    db.refresh(assignment)
    return assignment


def update_program_assignment(db: Session, assignment_id: int, data: ProgramAssignmentUpdate, organization_id: Optional[int] = None) -> LearningTrainingProgramAssignment:
    assignment = db.query(LearningTrainingProgramAssignment).filter(LearningTrainingProgramAssignment.id == assignment_id).first()
    if not assignment:
        raise NotFoundException("LearningTrainingProgramAssignment", assignment_id)
    update_data = data.model_dump(exclude_unset=True)
    for field, value in update_data.items():
        setattr(assignment, field, value)
    db.commit()
    db.refresh(assignment)
    return assignment


def remove_program_assignment(db: Session, assignment_id: int, organization_id: Optional[int] = None) -> None:
    assignment = db.query(LearningTrainingProgramAssignment).filter(LearningTrainingProgramAssignment.id == assignment_id).first()
    if not assignment:
        raise NotFoundException("LearningTrainingProgramAssignment", assignment_id)
    db.delete(assignment)
    db.commit()


def get_program_assignments(db: Session, program_id: int, organization_id: Optional[int] = None) -> list[dict]:
    get_training_program_by_id(db, program_id, organization_id)
    assignments = db.query(LearningTrainingProgramAssignment).filter(
        LearningTrainingProgramAssignment.program_id == program_id
    ).all()
    items = []
    for a in assignments:
        items.append({
            "id": a.id,
            "program_id": a.program_id,
            "employee_id": a.employee_id,
            "status": a.status,
            "attended_at": a.attended_at,
            "created_at": a.created_at,
            "employee_name": a.employee.full_name if a.employee else None,
        })
    return items


def create_calendar_event(db: Session, data: CalendarEventCreate, created_by: int, organization_id: Optional[int] = None) -> LearningCalendarEvent:
    event = LearningCalendarEvent(**data.model_dump(), created_by=created_by)
    if organization_id is not None:
        event.organization_id = organization_id
    db.add(event)
    db.commit()
    db.refresh(event)
    return event


def get_calendar_events(
    db: Session,
    start_date: Optional[date] = None,
    end_date: Optional[date] = None,
    organization_id: Optional[int] = None,
) -> list[LearningCalendarEvent]:
    query = db.query(LearningCalendarEvent)
    if start_date:
        query = query.filter(LearningCalendarEvent.event_date >= start_date)
    if end_date:
        query = query.filter(LearningCalendarEvent.event_date <= end_date)
    if organization_id is not None:
        query = query.filter(LearningCalendarEvent.organization_id == organization_id)
    return query.order_by(LearningCalendarEvent.event_date, LearningCalendarEvent.start_time).all()


def get_calendar_event_by_id(db: Session, event_id: int, organization_id: Optional[int] = None) -> LearningCalendarEvent:
    event = db.query(LearningCalendarEvent).filter(LearningCalendarEvent.id == event_id).first()
    if not event:
        raise NotFoundException("LearningCalendarEvent", event_id)
    return event


def update_calendar_event(db: Session, event_id: int, data: CalendarEventUpdate, organization_id: Optional[int] = None) -> LearningCalendarEvent:
    event = get_calendar_event_by_id(db, event_id, organization_id)
    update_data = data.model_dump(exclude_unset=True)
    for field, value in update_data.items():
        setattr(event, field, value)
    db.commit()
    db.refresh(event)
    return event


def delete_calendar_event(db: Session, event_id: int, organization_id: Optional[int] = None) -> None:
    event = get_calendar_event_by_id(db, event_id, organization_id)
    db.delete(event)
    db.commit()



def get_course_completion_report(db: Session, organization_id: Optional[int] = None) -> list[dict]:
    org_filter = [LearningCourse.organization_id == organization_id] if organization_id else []
    courses = db.query(LearningCourse).filter(*org_filter).all()
    result = []
    for course in courses:
        total = db.query(LearningEnrollment).filter(LearningEnrollment.course_id == course.id).count()
        completed = db.query(LearningEnrollment).filter(
            LearningEnrollment.course_id == course.id,
            LearningEnrollment.status == "completed",
        ).count()
        rate = round((completed / total * 100) if total > 0 else 0.0, 2)
        result.append({
            "course_id": course.id,
            "course_name": course.course_name,
            "total_enrollments": total,
            "completed": completed,
            "completion_rate": rate,
        })
    return result


def get_skill_gap_analysis(db: Session, organization_id: Optional[int] = None) -> list[dict]:
    from app.modules.hr.models import Employee, Department
    from sqlalchemy import func
    query = (
        db.query(
            Department.id.label("department_id"),
            Department.name.label("department_name"),
            LearningSkill.skill_name,
            func.avg(LearningSkill.proficiency_level).label("avg_level"),
            func.count(LearningSkill.id).label("employee_count"),
        )
        .join(Employee, LearningSkill.employee_id == Employee.id)
        .outerjoin(Department, Employee.department_id == Department.id)
    )
    if organization_id:
        query = query.filter(Employee.organization_id == organization_id)
    skills = (
        query.group_by(Department.id, Department.name, LearningSkill.skill_name)
        .all()
    )
    dept_map = {}
    for s in skills:
        dept_id = s.department_id or 0
        dept_name = s.department_name or "Unassigned"
        if dept_id not in dept_map:
            dept_map[dept_id] = {
                "department_id": dept_id,
                "department_name": dept_name,
                "skills": []
            }
        gap = max(0, 5 - (s.avg_level or 0))
        gap_level = "low" if gap <= 1 else "medium" if gap <= 2 else "high"
        dept_map[dept_id]["skills"].append({
            "skill_id": None,
            "skill_name": s.skill_name,
            "employee_count": s.employee_count,
            "gap": round(gap, 2),
            "gap_level": gap_level
        })
    return list(dept_map.values())


def get_learning_dashboard(db: Session, organization_id: Optional[int] = None) -> dict:
    org_filter = [LearningCourse.organization_id == organization_id] if organization_id else []
    total_courses = db.query(LearningCourse).filter(*org_filter).count()
    active_courses = db.query(LearningCourse).filter(LearningCourse.status == "active", *org_filter).count()

    enroll_org_filter = [LearningEnrollment.organization_id == organization_id] if organization_id else []
    total_enrollments = db.query(LearningEnrollment).filter(*enroll_org_filter).count()
    completed_enrollments = db.query(LearningEnrollment).filter(LearningEnrollment.status == "completed", *enroll_org_filter).count()
    completion_rate = round((completed_enrollments / total_enrollments * 100) if total_enrollments > 0 else 0.0, 2)

    cert_org_filter = [LearningCertification.organization_id == organization_id] if organization_id else []
    total_certifications = db.query(LearningCertification).filter(*cert_org_filter).count()

    skill_org_filter = [LearningSkill.organization_id == organization_id] if organization_id else []
    total_skills = db.query(LearningSkill).filter(*skill_org_filter).count()
    avg_skill_level = round(float(db.query(func.avg(LearningSkill.proficiency_level)).filter(*skill_org_filter).scalar() or 0.0), 2)

    attempt_org_filter = [LearningQuizAttempt.organization_id == organization_id] if organization_id else []
    in_progress_attempts = db.query(LearningQuizAttempt).filter(
        LearningQuizAttempt.status == "in_progress", *attempt_org_filter
    ).count()

    today = date.today()
    event_org_filter = [LearningCalendarEvent.organization_id == organization_id] if organization_id else []
    upcoming_events = db.query(LearningCalendarEvent).filter(
        LearningCalendarEvent.event_date >= today, *event_org_filter
    ).count()

    enrollment_trend = (
        db.query(
            func.to_char(LearningEnrollment.enrolled_at, "YYYY-MM").label("month"),
            func.count(LearningEnrollment.id).label("count"),
        )
        .filter(*enroll_org_filter)
        .group_by("month")
        .order_by("month")
        .all()
    )

    category_distribution = (
        db.query(
            LearningCourse.category,
            func.count(LearningCourse.id).label("count"),
        )
        .filter(LearningCourse.category.isnot(None), *org_filter)
        .group_by(LearningCourse.category)
        .all()
    )

    recent_enrollments_raw = (
        db.query(LearningEnrollment)
        .filter(*enroll_org_filter)
        .order_by(LearningEnrollment.enrolled_at.desc())
        .limit(10)
        .all()
    )

    recent_enrollments = []
    for e in recent_enrollments_raw:
        recent_enrollments.append({
            "id": e.id,
            "course_name": e.course.course_name if e.course else None,
            "employee_name": e.employee.full_name if e.employee else None,
            "status": e.status,
            "enrolled_at": str(e.enrolled_at) if e.enrolled_at else None,
        })

    return {
        "total_courses": total_courses,
        "active_courses": active_courses,
        "total_enrollments": total_enrollments,
        "completed_enrollments": completed_enrollments,
        "completion_rate": completion_rate,
        "total_certifications": total_certifications,
        "total_skills": total_skills,
        "avg_skill_level": avg_skill_level,
        "pending_assessments": in_progress_attempts,
        "upcoming_events": upcoming_events,
        "enrollment_trend": [{"month": m, "count": c} for m, c in enrollment_trend],
        "category_distribution": [{"category": cat, "count": cnt} for cat, cnt in category_distribution],
        "recent_enrollments": recent_enrollments,
    }


def export_course_completion_csv(db: Session, organization_id: Optional[int] = None) -> str:
    import io, csv
    data = get_course_completion_report(db, organization_id)
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(["Course ID", "Course Name", "Total Enrollments", "Completed Enrollments", "Completion Rate"])
    for row in data:
        writer.writerow([
            row["course_id"],
            row["course_name"],
            row["total_enrollments"],
            row["completed"],
            f"{row['completion_rate']}%"
        ])
    return output.getvalue()


def export_course_completion_excel(db: Session, organization_id: Optional[int] = None) -> bytes:
    import io
    from openpyxl import Workbook
    data = get_course_completion_report(db, organization_id)
    wb = Workbook()
    ws = wb.active
    ws.title = "Course Completion"
    ws.append(["Course ID", "Course Name", "Total Enrollments", "Completed Enrollments", "Completion Rate"])
    for row in data:
        ws.append([
            row["course_id"],
            row["course_name"],
            row["total_enrollments"],
            row["completed"],
            f"{row['completion_rate']}%"
        ])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def export_certifications_csv(db: Session, organization_id: Optional[int] = None) -> str:
    import io, csv
    certs = get_certifications(db, organization_id=organization_id)
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(["ID", "Employee ID", "Certification Name", "Issuing Organization", "Issue Date", "Expiry Date", "Credential URL", "Status"])
    for c in certs:
        writer.writerow([
            c.id,
            c.employee_id,
            c.certification_name,
            c.issuing_organization or "",
            c.issue_date,
            c.expiry_date or "",
            c.credential_url or "",
            c.status
        ])
    return output.getvalue()


def export_certifications_excel(db: Session, organization_id: Optional[int] = None) -> bytes:
    import io
    from openpyxl import Workbook
    certs = get_certifications(db, organization_id=organization_id)
    wb = Workbook()
    ws = wb.active
    ws.title = "Certifications"
    ws.append(["ID", "Employee ID", "Certification Name", "Issuing Organization", "Issue Date", "Expiry Date", "Credential URL", "Status"])
    for c in certs:
        ws.append([
            c.id,
            c.employee_id,
            c.certification_name,
            c.issuing_organization or "",
            str(c.issue_date),
            str(c.expiry_date) if c.expiry_date else "",
            c.credential_url or "",
            c.status
        ])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def export_skill_gap_csv(db: Session, organization_id: Optional[int] = None) -> str:
    import io, csv
    gaps = get_skill_gap_analysis(db, organization_id)
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(["Department Name", "Skill Name", "Employee Count", "Gap Level"])
    for dept in gaps:
        for s in dept["skills"]:
            writer.writerow([
                dept["department_name"],
                s["skill_name"],
                s["employee_count"],
                s["gap_level"]
            ])
    return output.getvalue()


def export_skill_gap_excel(db: Session, organization_id: Optional[int] = None) -> bytes:
    import io
    from openpyxl import Workbook
    gaps = get_skill_gap_analysis(db, organization_id)
    wb = Workbook()
    ws = wb.active
    ws.title = "Skill Gap Analysis"
    ws.append(["Department Name", "Skill Name", "Employee Count", "Gap Level"])
    for dept in gaps:
        for s in dept["skills"]:
            ws.append([
                dept["department_name"],
                s["skill_name"],
                s["employee_count"],
                s["gap_level"]
            ])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
