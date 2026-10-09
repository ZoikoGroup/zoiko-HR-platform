"""ZHR-38: employee import - re-importing updates people, new codes never collide, and
the Status column understands 'Probation'."""

import io
from datetime import date
from unittest.mock import patch

import openpyxl
import pytest
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.core.code_generation import generate_employee_code
from app.database import Base
from app.modules.employee import service
from app.modules.employee.models import Employee, EmployeeStatus, EmploymentType, UserRole
from app.modules.hr.models import Organization, OrganizationStatus

HEADERS = ["First Name", "Last Name", "Email", "Job Title", "Date of Joining", "Status", "Employment Type"]


@pytest.fixture
def db():
    engine = create_engine("sqlite:///:memory:", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(engine)
    s = sessionmaker(bind=engine, expire_on_commit=False)()
    s.add(Organization(id=1, name="Test Co", organization_name="Test Co", organization_code="TE", status=OrganizationStatus.ACTIVE))
    s.add(Organization(id=2, name="Other Co", organization_name="Other Co", organization_code="OT", status=OrganizationStatus.ACTIVE))
    s.commit()
    with patch("app.modules.employee.service._notify_email"):
        yield s
    s.close()
    engine.dispose()


def _xlsx(rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(HEADERS)
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _import(db, rows, org=1):
    return service.import_employees_from_file(db=db, file_bytes=_xlsx(rows), filename="people.xlsx",
                                              organization_id=org, current_user_id=None)


def _row(n, status="Active", etype="Full Time"):
    return [f"First{n}", f"Last{n}", f"person{n}@example.com", "Engineer", "2026-09-01", status, etype]


def test_codes_do_not_collide_after_an_employee_is_deleted(db):
    assert _import(db, [_row(i) for i in range(1, 6)])["created"] == 5
    codes = sorted(e.employee_code for e in db.query(Employee).all())
    assert codes == [f"TEE0000{i}" for i in range(1, 6)]
    # delete a MIDDLE employee: a count-based generator would now hand out an existing code
    db.delete(db.query(Employee).filter(Employee.employee_code == "TEE00002").one())
    db.commit()
    res = _import(db, [_row(i) for i in range(6, 9)])
    assert (res["created"], res["failed"]) == (3, 0), res["errors"]
    new = sorted(e.employee_code for e in db.query(Employee).filter(Employee.email.in_([f"person{i}@example.com" for i in (6, 7, 8)])))
    assert new == ["TEE00006", "TEE00007", "TEE00008"]
    assert len({e.employee_code for e in db.query(Employee).all()}) == db.query(Employee).count()


def test_generated_codes_are_per_organization_and_never_reused(db):
    first = generate_employee_code(db, 1)
    assert first == "TEE00001"
    db.add(Employee(email="a@example.com", hashed_password="x", employee_code="TEE00007", role=UserRole.EMPLOYEE, first_name="A",
                    last_name="B", job_title="t", employment_type=EmploymentType.FULL_TIME, status=EmployeeStatus.ACTIVE,
                    date_of_joining=date.today(), organization_id=1))
    db.commit()
    assert generate_employee_code(db, 1) == "TEE00008"  # highest + 1, not a count
    assert generate_employee_code(db, 2) == "OTE00001"


def test_reimporting_the_sheet_updates_existing_people_including_status(db):
    assert _import(db, [_row(1), _row(2)])["created"] == 2
    res = _import(db, [_row(1, status="Inactive"), _row(2, etype="Contract")])
    assert (res["created"], res["updated"], res["failed"]) == (0, 2, 0), res["errors"]
    one = db.query(Employee).filter(Employee.email == "person1@example.com").one()
    two = db.query(Employee).filter(Employee.email == "person2@example.com").one()
    assert one.status == EmployeeStatus.INACTIVE and one.is_active is False  # the visible status AND the access flag
    assert two.employment_type == EmploymentType.CONTRACT
    assert db.query(Employee).count() == 2  # no duplicates


def test_status_back_to_active_reactivates(db):
    _import(db, [_row(1, status="Inactive")])
    assert db.query(Employee).one().is_active is False
    _import(db, [_row(1, status="Active")])
    emp = db.query(Employee).one()
    assert emp.status == EmployeeStatus.ACTIVE and emp.is_active is True


def test_probation_in_the_status_column_is_accepted(db):
    res = _import(db, [_row(1, status="Probation", etype=None)])
    assert (res["created"], res["skipped"]) == (1, 0), res["errors"]
    emp = db.query(Employee).one()
    assert emp.status == EmployeeStatus.ACTIVE and emp.employment_type == EmploymentType.PROBATION


def test_an_explicit_employment_type_wins_over_probation_status(db):
    _import(db, [_row(1, status="Probation", etype="Contract")])
    assert db.query(Employee).one().employment_type == EmploymentType.CONTRACT


def test_unknown_status_error_is_readable_and_lists_real_options(db):
    res = _import(db, [_row(1, status="Wandering")])
    assert res["skipped"] == 1
    msg = res["errors"][0]["error"]
    assert "Wandering" in msg and "active" in msg and "probation" in msg.lower()


def test_database_errors_are_reported_without_sql_dumps(db):
    with patch("app.core.code_generation.generate_employee_code", side_effect=Exception(
            '(psycopg2.errors.UniqueViolation) duplicate key value violates unique constraint "employees_employee_code_key"\n'
            'DETAIL: Key (employee_code)=(TEE00024) already exists.\n[SQL: INSERT INTO employees ...]'), create=True):
        res = _import(db, [_row(1)])
    assert res["failed"] == 1
    msg = res["errors"][0]["error"]
    assert "[SQL" not in msg and "psycopg2" not in msg


# ── the import reads the headings people actually use, and explains a file it cannot read
def _book(sheets):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for title, rows in sheets:
        ws = wb.create_sheet(title)
        for r in rows:
            ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _run(db, data, name="people.xlsx"):
    return service.import_employees_from_file(db=db, file_bytes=data, filename=name, organization_id=1, current_user_id=None)


def test_headings_with_stars_hints_and_other_wordings_are_understood(db):
    headers = ["First Name *", "Surname", "E-mail", "Position", "Date of Joining (YYYY-MM-DD)", "Mobile Number"]
    data = _book([("Staff", [headers, ["Asha", "Rao", "asha.rao@example.com", "Analyst", "2026-09-01", "9876543210"],
                             ["Ravi", "Iyer", "ravi.iyer@example.com", "Designer", "2026-09-02", "9876543211"]])])
    res = _run(db, data)
    assert (res["created"], res["failed"], res["skipped"]) == (2, 0, 0), res["errors"]
    asha = db.query(Employee).filter_by(email="asha.rao@example.com").one()
    assert (asha.first_name, asha.last_name, asha.job_title) == ("Asha", "Rao", "Analyst")


def test_a_title_row_or_notes_above_the_headings_do_not_stop_the_import_and_row_numbers_match_the_file(db):
    rows = [["Employees to import"], ["Fill one row per person"], [], ["First Name", "Last Name", "Email", "Job Title", "Date of Joining"],
            ["Asha", "Rao", "asha.rao@example.com", "Analyst", "2026-09-01"], ["", "", "", "", ""], ["Bad", "Row", "not-an-email", "Analyst", "2026-09-01"]]
    res = _run(db, _book([("Staff", rows)]))
    assert res["created"] == 1
    assert [e["row"] for e in res["errors"]] == [7], "the bad row is reported at the row it is on in the spreadsheet"


def test_the_sheet_that_looks_like_a_staff_list_is_used_even_when_it_is_not_the_first(db):
    data = _book([("Read me", [["How to use this file"], ["Fill in the Staff sheet"]]),
                  ("Staff", [["First Name", "Last Name", "Email", "Job Title", "Date of Joining"], ["Asha", "Rao", "asha.rao@example.com", "Analyst", "2026-09-01"]])])
    assert _run(db, data)["created"] == 1


def test_a_file_without_the_required_columns_gets_one_clear_message_not_one_per_row(db):
    data = _book([("Staff", [["Name", "Mail", "Team"]] + [[f"P{i}", f"p{i}@example.com", "Ops"] for i in range(1, 21)])])
    res = _run(db, data)
    assert res["created"] == 0 and len(res["errors"]) == 1, res["errors"]
    message = res["errors"][0]["error"]
    assert "no column for" in message and "First Name" in message and "Job Title" in message and "Date of Joining" in message
    assert "Columns found: Name, Mail, Team" in message and "template" in message.lower()


def test_csv_with_semicolons_a_bom_and_starred_headings_imports(db):
    text = "\ufeffFirst Name*;Last Name*;Email*;Job Title*;Date of Joining*\nAsha;Rao;asha.rao@example.com;Analyst;2026-09-01\nRavi;Iyer;ravi.iyer@example.com;Designer;01/09/2026\n"
    res = _run(db, text.encode("utf-8"), "people.csv")
    assert (res["created"], res["failed"], res["skipped"]) == (2, 0, 0), res["errors"]


def test_the_template_file_itself_still_imports(db):
    with open("../frontend/public/templates/employee-import-template.xlsx", "rb") as f:
        res = _run(db, f.read(), "employee-import-template.xlsx")
    assert res["errors"] == [] or all("Missing required" in e["error"] or "date" in e["error"].lower() for e in res["errors"]), res["errors"]
