"""Country-localised employee import templates: headers, sample row and currency follow ?country=."""
from io import BytesIO

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.database import Base, get_db
from app.core.dependencies import get_current_user
from app.modules.employee.router import employee_router
from app.modules.employee.service import _generate_import_template_bytes, _normalise_header


@pytest.fixture
def client():
    engine = create_engine("sqlite:///:memory:", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(engine)
    s = sessionmaker(bind=engine, expire_on_commit=False)()
    app = FastAPI()
    app.include_router(employee_router)

    def _override_db():
        yield s

    app.dependency_overrides[get_db] = _override_db
    app.dependency_overrides[get_current_user] = lambda: type("U", (), {"id": 1, "organization_id": 1, "role": "admin"})()
    with TestClient(app) as c:
        yield c
    s.close()
    engine.dispose()


def _headers(blob):
    import openpyxl
    wb = openpyxl.load_workbook(BytesIO(blob))
    ws = wb.active
    return [c.value for c in ws[1]], [c.value for c in ws[2]]


def test_default_template_is_india_flavoured(client):
    r = client.get("/hr/employee-management/employees/import/template")
    assert r.status_code == 200
    assert "employee-import-template-default.xlsx" in r.headers["content-disposition"]
    headers, sample = _headers(r.content)
    assert "PAN Number" in headers and "IFSC Code" in headers
    assert "Currency" in headers
    assert sample[headers.index("Country")] == "India"
    assert sample[headers.index("Currency")] == "INR"


def test_united_states_template(client):
    r = client.get("/hr/employee-management/employees/import/template", params={"country": "United States"})
    assert r.status_code == 200
    assert "employee-import-template-united-states.xlsx" in r.headers["content-disposition"]
    headers, sample = _headers(r.content)
    assert "SSN" in headers and "Routing Number" in headers
    assert "PAN Number" not in headers and "IFSC Code" not in headers
    assert sample[headers.index("Country")] == "United States"
    assert sample[headers.index("Currency")] == "USD"


def test_unknown_country_gets_international_headings_not_indian_ones():
    headers, sample = _headers(_generate_import_template_bytes(country="Atlantis"))
    assert "PAN Number" not in headers and "Tax ID" in headers and "SWIFT Code" in headers
    assert sample[headers.index("Country")] == "Atlantis"


def test_no_country_keeps_the_india_template():
    headers, sample = _headers(_generate_import_template_bytes())
    assert "PAN Number" in headers and sample[headers.index("Country")] == "India"


def test_localised_headers_parse_back_to_known_fields():
    assert _normalise_header("SSN") == "pan_number"
    assert _normalise_header("Routing Number") == "bank_ifsc"
    assert _normalise_header("Sort Code") == "bank_ifsc"
    assert _normalise_header("National Insurance Number") == "pan_number"
    assert _normalise_header("Currency") == "currency"
    assert _normalise_header("PAN Number") == "pan_number"


from app.modules.employee import import_countries  # noqa: E402


@pytest.mark.parametrize("country", [c["name"] for c in import_countries.COUNTRIES])
def test_every_country_template_reads_back_into_the_right_fields(country):
    """Fill nothing in: the sample row of each country's template must import into the right employee fields."""
    from app.modules.employee.service import _parse_file
    cfg = import_countries.for_country(country)
    blob = _generate_import_template_bytes(country=country)
    headers, sample = _headers(blob)
    assert sample[headers.index("Country")] == country
    assert sample[headers.index("Currency")] == cfg["currency"]
    result = {"errors": [], "failed": 0, "created": 0, "updated": 0, "total_rows": 0}
    rows = _parse_file(blob, "template.xlsx", result)
    assert len(rows) == 1, result
    row = {_normalise_header(k): v for k, v in rows[0].items() if k != "__row"}      # as the import maps headings
    for field in import_countries.FIELDS:
        header, _label, example = cfg["ids"][field]
        assert _normalise_header(header) == field, (country, header)
        if example:
            assert str(row.get(field)) == example, (country, field, row.get(field))
    assert row.get("country") == country and row.get("currency") == cfg["currency"]
    assert row.get("city") == cfg["place"][1]
    if cfg["place"][3]:
        assert str(row.get("pincode")) == cfg["place"][3], (country, row.get("pincode"))
    if cfg["place"][2]:
        assert row.get("state") == cfg["place"][2], (country, row.get("state"))


def test_twenty_countries_each_with_their_own_currency_and_no_heading_clash():
    names = [c["name"] for c in import_countries.COUNTRIES]
    assert len(names) == 20 and len(set(names)) == 20
    from app.core import regions
    for c in import_countries.COUNTRIES:
        assert regions.currency_for(c["name"]) == c["currency"], c["name"]      # matches what User Management shows


def test_template_countries_endpoint(client):
    rows = client.get("/hr/employee-management/employees/import/template-countries").json()
    assert len(rows) == 20
    de = next(r for r in rows if r["name"] == "Germany")
    assert de["currency"] == "EUR" and de["identifiers"][0] == "Tax ID (Steuer-ID)"
