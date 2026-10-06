"""ZHR-56: attendance CSV and Excel exports are two different, correct, organization-scoped formats."""

import csv
import io
from datetime import date, datetime, timedelta

import openpyxl
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.core.dependencies import get_current_user
from app.database import Base, get_db
from app.modules.employee.models import Employee, EmployeeStatus, EmploymentType, UserRole
from app.modules.hr.models import AttendanceRecord, AttendanceStatus, Department, Organization, OrganizationStatus


@pytest.fixture
def client():
    from app.main import app

    engine = create_engine("sqlite:///:memory:", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(engine)
    s = sessionmaker(bind=engine, expire_on_commit=False)()
    s.add_all([Organization(id=1, name="Acme", status=OrganizationStatus.ACTIVE), Organization(id=2, name="Globex", status=OrganizationStatus.ACTIVE)])
    s.commit()
    s.add(Department(id=1, name="Engineering", code="ENG", organization_id=1))
    s.commit()
    people = {}
    for key, role, org, first, last in (("admin", UserRole.ADMIN, 1, "Priya", "Shah"), ("emp", UserRole.EMPLOYEE, 1, "=cmd", "Évelyne"),
                                        ("other", UserRole.EMPLOYEE, 2, "Other", "Org")):
        e = Employee(email=f"{key}@example.com", hashed_password="x", employee_code=f"C-{key}", role=role, first_name=first, last_name=last,
                     job_title="t", employment_type=EmploymentType.FULL_TIME, status=EmployeeStatus.ACTIVE, date_of_joining=date.today(),
                     organization_id=org, department_id=1 if org == 1 else None)
        s.add(e)
        people[key] = e
    s.commit()
    today = date.today()
    d = datetime.combine(today, datetime.min.time())
    s.add_all([
        AttendanceRecord(employee_id=people["emp"].id, organization_id=1, date=today, status=AttendanceStatus.PRESENT,
                         check_in=d.replace(hour=9, minute=5), check_out=d.replace(hour=17, minute=30), total_hours=8.42, notes="@sum"),
        AttendanceRecord(employee_id=people["admin"].id, organization_id=1, date=today - timedelta(days=1), status=AttendanceStatus.ABSENT),
        AttendanceRecord(employee_id=people["emp"].id, organization_id=1, date=today - timedelta(days=2), status=AttendanceStatus.PRESENT, is_deleted=True),
        AttendanceRecord(employee_id=people["other"].id, organization_id=2, date=today, status=AttendanceStatus.PRESENT),
    ])
    s.commit()
    box = {"user": people["admin"]}
    app.dependency_overrides[get_db] = lambda: s
    app.dependency_overrides[get_current_user] = lambda: box["user"]
    c = TestClient(app, raise_server_exceptions=False)
    c.box, c.people, c.today = box, people, today
    yield c
    app.dependency_overrides.pop(get_db, None)
    app.dependency_overrides.pop(get_current_user, None)
    s.close()
    engine.dispose()


def _csv_rows(resp):
    text = resp.content.decode("utf-8-sig")
    return list(csv.reader(io.StringIO(text)))


def _xlsx(resp):
    return openpyxl.load_workbook(io.BytesIO(resp.content))


def test_csv_and_excel_are_genuinely_different_formats(client):
    csv_resp = client.get("/hr/attendance/export/csv")
    xlsx_resp = client.get("/hr/attendance/export/excel")
    assert csv_resp.status_code == 200 and xlsx_resp.status_code == 200, (csv_resp.text, xlsx_resp.text)
    assert csv_resp.headers["content-type"].startswith("text/csv")
    assert xlsx_resp.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert csv_resp.headers["content-disposition"].endswith('.csv"')
    assert xlsx_resp.headers["content-disposition"].endswith('.xlsx"')
    assert xlsx_resp.content[:2] == b"PK"  # a zip container, i.e. a real workbook
    assert not csv_resp.content.startswith(b"PK")
    csv_resp.content.decode("utf-8")  # plain text, readable as such


def test_csv_content(client):
    rows = _csv_rows(client.get("/hr/attendance/export/csv"))
    assert rows[0] == ["Employee Code", "Employee Name", "Department", "Date", "Check In", "Check Out", "Status", "Hours Worked", "Notes"]
    assert len(rows) == 3  # header + this organization's 2 live records (not the deleted one, not the other org's)
    present = next(r for r in rows[1:] if r[6] == "present")
    assert present[0] == "C-emp" and present[2] == "Engineering" and present[3] == client.today.isoformat()
    assert (present[4], present[5], present[7]) == ("09:05", "17:30", "8.42")
    assert "Évelyne" in present[1]  # accents survive (UTF-8 with BOM for Excel)
    assert present[1].startswith("'=cmd")  # a name that looks like a formula is neutralised
    assert present[8] == "'@sum"


def test_excel_content_has_typed_cells_a_styled_header_and_a_summary(client):
    wb = _xlsx(client.get("/hr/attendance/export/excel"))
    ws = wb["Attendance"]
    assert [c.value for c in ws[1]] == ["Employee Code", "Employee Name", "Department", "Date", "Check In", "Check Out", "Status", "Hours Worked", "Notes"]
    assert ws[1][0].font.bold and ws.freeze_panes == "A2" and ws.auto_filter.ref
    assert ws.max_row == 3
    row = next(r for r in ws.iter_rows(min_row=2) if r[6].value == "Present")
    assert isinstance(row[3].value, (date, datetime))  # a real date cell
    assert row[3].number_format == "DD-MMM-YYYY"
    assert row[4].value.hour == 9 and row[5].value.minute == 30  # real time cells
    assert row[7].value == pytest.approx(8.42)
    summary = {r[0].value: r[1].value for r in wb["Summary"].iter_rows(min_row=2)}
    assert summary == {"Absent": 1, "Present": 1, "Total": 2}


@pytest.mark.parametrize("kind", ["csv", "excel"])
def test_only_this_organizations_attendance_is_exported(client, kind):
    resp = client.get(f"/hr/attendance/export/{kind}")
    body = resp.content.decode("utf-8", "ignore") if kind == "csv" else " ".join(str(c.value) for r in _xlsx(resp)["Attendance"].iter_rows() for c in r)
    assert "C-other" not in body and "Other" not in body


@pytest.mark.parametrize("kind", ["csv", "excel"])
def test_date_range_and_employee_filters(client, kind):
    day = (client.today - timedelta(days=1)).isoformat()
    resp = client.get(f"/hr/attendance/export/{kind}", params={"date_from": day, "date_to": day})
    count = len(_csv_rows(resp)) - 1 if kind == "csv" else _xlsx(resp)["Attendance"].max_row - 1
    assert count == 1
    resp = client.get(f"/hr/attendance/export/{kind}", params={"employee_id": client.people["emp"].id})
    count = len(_csv_rows(resp)) - 1 if kind == "csv" else _xlsx(resp)["Attendance"].max_row - 1
    assert count == 1


@pytest.mark.parametrize("kind", ["csv", "excel"])
def test_an_empty_export_is_still_a_valid_file(client, kind):
    resp = client.get(f"/hr/attendance/export/{kind}", params={"date_from": "2001-01-01", "date_to": "2001-01-02"})
    assert resp.status_code == 200
    if kind == "csv":
        assert len(_csv_rows(resp)) == 1
    else:
        assert _xlsx(resp)["Attendance"].max_row == 1


@pytest.mark.parametrize("kind", ["csv", "excel"])
def test_employees_cannot_export_everyone_s_attendance(client, kind):
    client.box["user"] = client.people["emp"]
    assert client.get(f"/hr/attendance/export/{kind}").status_code == 403


def test_hours_are_derived_from_check_in_and_out_when_no_total_is_stored(client):
    from app.database import get_db

    session = client.app.dependency_overrides[get_db]()
    d = datetime.combine(client.today - timedelta(days=5), datetime.min.time())
    session.add(AttendanceRecord(employee_id=client.people["admin"].id, organization_id=1, date=d.date(), status=AttendanceStatus.PRESENT,
                                 check_in=d.replace(hour=9), check_out=d.replace(hour=17, minute=30)))
    session.commit()
    rows = _csv_rows(client.get("/hr/attendance/export/csv", params={"date_from": d.date().isoformat(), "date_to": d.date().isoformat()}))
    assert rows[1][7] == "8.50"
